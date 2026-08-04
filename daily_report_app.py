# -*- coding: utf-8 -*-
"""
Daily Report App v2 -- PT. Garuda Prima Aksara
Run:  python daily_report_app.py
Open: http://localhost:5050
"""

# ============================================================
#  PREREQUISITE CHECKER  -- runs before anything else
#  Installs missing packages automatically so the app always
#  starts cleanly on any machine.
# ============================================================
import sys, subprocess

_REQUIRED = {
    "flask":      "flask",
    "reportlab":  "reportlab",
    "PIL":        "pillow",
    "docx":       "python-docx",
    "anthropic":  "anthropic",
}

def _check_and_install():
    missing = []
    for mod, pkg in _REQUIRED.items():
        try:
            __import__(mod)
        except ImportError:
            missing.append(pkg)

    if not missing:
        return  # all good

    print()
    print("  ============================================")
    print("   Installing missing packages, please wait...")
    print("  ============================================")
    for pkg in missing:
        print(f"  Installing {pkg}...", flush=True)
        result = subprocess.run(
            [sys.executable, "-m", "pip", "install", pkg,
             "--quiet", "--no-warn-script-location"],
            capture_output=True, text=True
        )
        if result.returncode != 0:
            print(f"  ERROR installing {pkg}:")
            print(result.stderr[:500])
            print()
            print("  Fix: right-click the .bat file and choose 'Run as administrator'")
            print("  Then close and reopen the app.")
            input("  Press Enter to exit...")
            sys.exit(1)
        print(f"  {pkg} installed OK")

    # Verify everything is now importable
    still_missing = []
    for mod, pkg in _REQUIRED.items():
        try:
            __import__(mod)
        except ImportError:
            still_missing.append(pkg)

    if still_missing:
        print()
        print(f"  ERROR: Could not import: {still_missing}")
        print("  Try running as administrator, or check internet connection.")
        input("  Press Enter to exit...")
        sys.exit(1)

    print("  All packages ready!")
    print()

_check_and_install()

# ============================================================
#  Now safe to import everything
# ============================================================
import os, json, io, base64, uuid, re, string, copy, math
from datetime import datetime
from flask import Flask, render_template, request, jsonify, send_file, session, redirect, url_for
from PIL import Image, ImageOps, UnidentifiedImageError

# ── PDF engine ────────────────────────────────────────────────────────────────
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
    TableStyle, CondPageBreak, KeepTogether, Image as RLImage)
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.lib.utils import ImageReader

W, H = A4
M = 15 * mm
CW = W - 2 * M
PDF_ASSET_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'static', 'pdf_assets')
BUNDLED_LOGOS = {
    'gpa': os.path.join(PDF_ASSET_DIR, 'gpa_logo.png'),
    'kn': os.path.join(PDF_ASSET_DIR, 'kn_logo.png'),
}
BUNDLED_HEADER_LOGOS = {
    'gpa': os.path.join(PDF_ASSET_DIR, 'gpa_logo_header_text_plus_1pt.png'),
    'kn': os.path.join(PDF_ASSET_DIR, 'kn_logo_header.png'),
}

def _is_drawable_logo(path):
    if not path or not os.path.isfile(path):
        return False
    try:
        ImageReader(path).getSize()
        return True
    except Exception:
        return False

def resolve_logo_path(configured_path, which):
    """Prefer a valid uploaded raster logo, otherwise use the bundled logo."""
    if _is_drawable_logo(configured_path):
        return configured_path
    fallback = BUNDLED_LOGOS.get(which, '')
    return fallback if _is_drawable_logo(fallback) else ''

_SS = getSampleStyleSheet()
def S(nm, **kw): return ParagraphStyle('_', parent=_SS.get(nm, _SS['Normal']), **kw)

def make_styles(cfg):
    t = cfg.get('theme', {})
    PRI   = colors.HexColor(t.get('primary',   '#003366'))
    SEC   = colors.HexColor(t.get('secondary', '#005B99'))
    ACC   = colors.HexColor(t.get('accent',    '#C89010'))
    AREA  = colors.HexColor(t.get('area_hdr',  '#1A5276'))
    LB    = colors.HexColor(t.get('light_bg',  '#D6E8F7'))
    return dict(
        PRI=PRI, SEC=SEC, ACC=ACC, AREA=AREA, LB=LB,
        sec_s  = S('Normal', fontSize=10.5, textColor=colors.white, fontName='Helvetica-Bold',
                    leading=12.5, spaceBefore=0, spaceAfter=0),
        area_s = S('Normal', fontSize=9, textColor=colors.white, fontName='Helvetica-Bold',
                    leading=10.5, spaceBefore=0, spaceAfter=0),
        sub_s  = S('Normal', fontSize=8.3, textColor=PRI, fontName='Helvetica-Bold', spaceBefore=2, spaceAfter=1),
        body_s = S('Normal', fontSize=8,   leading=10, spaceAfter=1),
        sm_s   = S('Normal', fontSize=7.5, leading=9, spaceAfter=1),
        tbl_s  = S('Normal', fontSize=7.5, leading=9, spaceBefore=0, spaceAfter=0,
                    splitLongWords=1),
        tbl_c_s= S('Normal', fontSize=7.5, leading=9, spaceBefore=0, spaceAfter=0,
                    splitLongWords=1, alignment=1),
        wx_h_s = S('Normal', fontSize=7, leading=8, textColor=colors.white,
                    fontName='Helvetica-Bold', spaceBefore=0, spaceAfter=0,
                    splitLongWords=1, alignment=1),
        bold_s = S('Normal', fontSize=8,   fontName='Helvetica-Bold', leading=10),
        ital_s = S('Normal', fontSize=7.5, fontName='Helvetica-Oblique', leading=9,
                    textColor=colors.HexColor('#555555')),
    )

def base_ts(extra=None):
    b = [('BACKGROUND',(0,0),(-1,0),colors.HexColor('#003366')),
         ('TEXTCOLOR',(0,0),(-1,0),colors.white),
         ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('FONTSIZE',(0,0),(-1,-1),7.5),
         ('FONTNAME',(0,1),(-1,-1),'Helvetica'),('GRID',(0,0),(-1,-1),0.4,colors.HexColor('#CCCCCC')),
         ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F2F2F2')]),
         ('VALIGN',(0,0),(-1,-1),'TOP'),('ALIGN',(0,0),(-1,0),'CENTER'),
         ('TOPPADDING',(0,0),(-1,-1),1.5),('BOTTOMPADDING',(0,0),(-1,-1),1.5),
         ('LEFTPADDING',(0,0),(-1,-1),3),('RIGHTPADDING',(0,0),(-1,-1),3)]
    if extra: b.extend(extra)
    return TableStyle(b)

class _NC(rl_canvas.Canvas):
    _meta = {}
    def __init__(self, *a, **kw):
        super().__init__(*a, **kw); self._saved_page_states = []
    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__)); self._startPage()
    def save(self):
        n = len(self._saved_page_states)
        for s in self._saved_page_states:
            self.__dict__.update(s); self._hf(n); super().showPage()
        super().save()

    def _draw_logo_badge(self, path, x, y, box_w, box_h, show_card=True):
        """Draw a supplied logo without distortion, optionally on a white card."""
        try:
            reader = ImageReader(path)
            image_w, image_h = reader.getSize()
            # Keep the artwork at the exact same size it had inside the former
            # white card; hiding the card must not enlarge either logo.
            pad = 0.55 * mm
            scale = min((box_w - 2*pad) / image_w, (box_h - 2*pad) / image_h)
            draw_w = image_w * scale
            draw_h = image_h * scale
            if show_card:
                self.setFillColor(colors.white)
                self.setStrokeColor(colors.HexColor('#D5DCE4'))
                self.setLineWidth(0.25)
                self.roundRect(x, y, box_w, box_h, 0.45*mm, fill=1, stroke=1)
            self.drawImage(
                reader,
                x + (box_w - draw_w) / 2,
                y + (box_h - draw_h) / 2,
                width=draw_w,
                height=draw_h,
                preserveAspectRatio=True,
                mask='auto',
            )
            return True
        except Exception:
            return False

    def _draw_fitted_center(self, text, x_min, x_max, y, font, max_size, min_size):
        """Fit one metadata line between both logo cards without overlap."""
        rendered = ' '.join(str(text or '').split())
        available = max(1, x_max - x_min)
        size = max_size
        while size > min_size and self.stringWidth(rendered, font, size) > available:
            size = max(min_size, size - 0.25)
        if self.stringWidth(rendered, font, size) > available:
            suffix = '...'
            while rendered and self.stringWidth(rendered + suffix, font, size) > available:
                rendered = rendered[:-1]
            rendered = rendered.rstrip() + suffix
        self.setFont(font, size)
        self.drawCentredString((x_min + x_max) / 2, y, rendered)

    def _draw_wrapped_center(self, text, x_min, x_max, center_y, font,
                             max_size, min_size, leading, max_lines=2):
        """Wrap a long header value while keeping it clear of both logo cards."""
        rendered = ' '.join(str(text or '').split())
        available = max(1, x_max - x_min)

        def wrap_words(font_size):
            lines = []
            current = ''
            for word in rendered.split():
                candidate = f'{current} {word}'.strip()
                if not current or self.stringWidth(candidate, font, font_size) <= available:
                    current = candidate
                else:
                    lines.append(current)
                    current = word
            if current:
                lines.append(current)
            return lines or ['']

        size = max_size
        lines = wrap_words(size)
        while len(lines) > max_lines and size > min_size:
            size = max(min_size, size - 0.25)
            lines = wrap_words(size)

        if len(lines) > max_lines:
            lines = lines[:max_lines - 1] + [' '.join(lines[max_lines - 1:])]
            suffix = '...'
            last = lines[-1]
            while last and self.stringWidth(last + suffix, font, size) > available:
                last = last[:-1]
            lines[-1] = last.rstrip() + suffix

        self.setFont(font, size)
        first_y = center_y + ((len(lines) - 1) * leading / 2)
        for index, line in enumerate(lines):
            self.drawCentredString(
                (x_min + x_max) / 2,
                first_y - index * leading,
                line,
            )

    def _hf(self, n):
        m = self.__class__._meta
        date=m.get('date',''); day=m.get('day',''); proj=m.get('proj','')
        loc=m.get('loc',''); cust=m.get('cust','')
        t = m.get('theme', {})
        PRI = colors.HexColor(t.get('primary','#003366'))
        ACC = colors.HexColor(t.get('accent','#C89010'))
        self.saveState()
        header_h = 36 * mm
        accent_h = 2 * mm
        logo_side_margin = 11 * mm
        self.setFillColor(PRI); self.rect(0, H-header_h, W, header_h, fill=1, stroke=0)
        # Logos + configurable header text
        logo_gpa    = m.get('logo_gpa','')
        logo_kn     = m.get('logo_kn','')
        show_gpa    = m.get('show_logo_gpa', True)
        show_kn     = m.get('show_logo_kn',  True)
        gpa_card    = m.get('logo_gpa_card', True)
        kn_card     = m.get('logo_kn_card', True)
        gpa_w       = m.get('logo_gpa_w', 28) * mm
        gpa_h       = m.get('logo_gpa_h', 12) * mm
        gpa_yo      = m.get('logo_gpa_y_off', 0) * mm
        kn_w        = m.get('logo_kn_w',  28) * mm
        kn_h        = m.get('logo_kn_h',  12) * mm
        kn_yo       = m.get('logo_kn_y_off',  0) * mm
        co_name     = m.get('company_name',  'PT. GARUDA PRIMA AKSARA')
        proj_title  = m.get('project_title', 'Electrical Installation & Construction')
        gpa_vis = show_gpa and _is_drawable_logo(logo_gpa)
        kn_vis  = show_kn  and _is_drawable_logo(logo_kn)
        # Lower both marks slightly so their visual centre sits more naturally
        # within the full header composition without changing their size.
        badge_center_y = H - 15*mm
        if gpa_vis:
            gpa_vis = self._draw_logo_badge(
                logo_gpa, logo_side_margin, badge_center_y - gpa_h/2 + gpa_yo, gpa_w, gpa_h,
                show_card=gpa_card)
        if kn_vis:
            kn_vis = self._draw_logo_badge(
                logo_kn, W-logo_side_margin-kn_w, badge_center_y - kn_h/2 + kn_yo, kn_w, kn_h,
                show_card=kn_card)

        text_left = logo_side_margin + (gpa_w + 3*mm if gpa_vis else 0)
        text_right = W - logo_side_margin - (kn_w + 3*mm if kn_vis else 0)
        # Keep every centre-column line on the true page centre. Because the
        # GPA and KN marks have different widths, use the narrower clearance
        # on both sides instead of centring inside an asymmetric gap.
        safe_half_width = min(W/2 - text_left, text_right - W/2)
        text_left = W/2 - safe_half_width
        text_right = W/2 + safe_half_width
        self.setFillColor(colors.white)
        self._draw_fitted_center(co_name, text_left, text_right, H-7*mm,
                                 'Helvetica-Bold', 10, 6.5)
        self._draw_fitted_center(
            f'Daily Activity Report  |  {cust}',
            text_left, text_right, H-11.5*mm, 'Helvetica', 7.3, 5.2)
        self._draw_wrapped_center(
            proj_title, text_left, text_right, H-17.5*mm,
            'Helvetica-Bold', 6.7, 5.2, 3.5*mm, max_lines=2)
        self._draw_fitted_center(
            f'Date: {date}  |  Day: {day}  |  Project: {proj}',
            M, W-M, H-24.5*mm, 'Helvetica', 6.7, 5.2)

        # Keep report metadata visible even when the KN logo is present.
        self.setFont('Helvetica', 6.8)
        self.drawString(logo_side_margin, H-31.5*mm, f'LOCATION: {loc}')
        self.drawCentredString(W/2, H-31.5*mm, f'CUSTOMER: {cust}')
        self.drawRightString(W-logo_side_margin, H-31.5*mm, f'DAY {day}')

        self.setFillColor(ACC); self.rect(0, H-header_h-accent_h, W, accent_h, fill=1, stroke=0)
        self.setFillColor(PRI); self.rect(0, 0, W, 8*mm, fill=1, stroke=0)
        self.setFillColor(colors.white); self.setFont('Helvetica',6.5)
        self.drawString(M, 2.7*mm, 'PT. Garuda Prima Aksara  |  Confidential')
        self.drawCentredString(W/2, 2.7*mm, f'Daily Activity Report  |  {date}  |  PT. KN')
        self.drawRightString(W-M, 2.7*mm, f'Page {self._pageNumber} of {n}')
        self.restoreState()

def _esc(s):
    """Escape HTML entities so ReportLab Paragraph never sees raw & < > from user input."""
    return str(s).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

PDF_PHOTO_CAPTION_MAX_CHARS = 300
PDF_PHOTO_AREA_MAX_CHARS = 120

OVERALL_PROGRESS_FIELDS = (
    'description',
    'duration',
    'weight_factor',
    'start',
    'finish',
    'cumulative_previous_plan',
    'cumulative_previous_actual',
    'this_period_plan',
    'this_period_actual',
    'cumulative_to_date_plan',
    'cumulative_to_date_actual',
    'deviation',
)

OVERALL_PROGRESS_PERCENT_FIELDS = (
    'cumulative_previous_plan',
    'cumulative_previous_actual',
    'this_period_plan',
    'this_period_actual',
    'cumulative_to_date_plan',
    'cumulative_to_date_actual',
)

def _bounded_pdf_text(value, max_chars):
    """Keep user text inside non-splittable PDF photo-card rows."""
    text = str(value or '').strip()
    if len(text) <= max_chars:
        return text
    return text[:max_chars - 3].rstrip() + '...'

def _coerce_bool(value, default=True):
    """Parse report flags without treating the string ``false`` as enabled."""
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    if isinstance(value, str):
        normalised = value.strip().lower()
        if normalised in {'true', '1', 'yes', 'on'}:
            return True
        if normalised in {'false', '0', 'no', 'off'}:
            return False
    return default

def _normalise_overall_progress(rows):
    """Return only usable progress rows while keeping old drafts compatible."""
    if not isinstance(rows, list):
        return []
    normalised = []
    for item in rows:
        if not isinstance(item, dict):
            continue
        row = {}
        for field in OVERALL_PROGRESS_FIELDS:
            value = item.get(field, '')
            row[field] = str('' if value is None else value).strip()
        if any(row.values()):
            normalised.append(row)
    return normalised

def _progress_number(value):
    """Parse a flexible percentage value such as 95,25%, 95.25, or 95."""
    if value is None or isinstance(value, bool):
        return None
    text = str(value).strip().replace('%', '').replace(' ', '')
    if not text:
        return None
    if ',' in text and '.' not in text:
        text = text.replace(',', '.')
    else:
        text = text.replace(',', '')
    try:
        number = float(text)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None

def _progress_percent_text(value):
    """Display numeric progress consistently while preserving non-numeric text."""
    text = str('' if value is None else value).strip()
    if not text:
        return ''
    number = _progress_number(text)
    if number is None:
        return text
    if text.endswith('%'):
        return text
    return f'{number:g}%'

def _overall_progress_totals(rows):
    """Calculate the weighted overall values shown in the summary row."""
    totals = {}
    for field in OVERALL_PROGRESS_PERCENT_FIELDS:
        total = 0.0
        found = False
        for row in rows:
            weight = _progress_number(row.get('weight_factor'))
            value = _progress_number(row.get(field))
            if weight is None or value is None:
                continue
            total += weight * value / 100.0
            found = True
        totals[field] = total if found else None

    plan = totals.get('cumulative_to_date_plan')
    actual = totals.get('cumulative_to_date_actual')
    totals['deviation'] = actual - plan if plan is not None and actual is not None else None
    return totals

def _group_report_photos(areas, per_row=3):
    """Build independent photo-grid rows for each report area."""
    if not isinstance(areas, list) or per_row < 1:
        return []
    grouped_areas = []
    for area in areas:
        if not isinstance(area, dict):
            continue
        area_id = area.get('id', '')
        photos = area.get('photos', [])
        if not isinstance(photos, list):
            continue
        entries = [(area_id, photo) for photo in photos if isinstance(photo, dict)]
        if not entries:
            continue
        rows = [entries[index:index + per_row] for index in range(0, len(entries), per_row)]
        grouped_areas.append((area_id, rows))
    return grouped_areas

def _prepare_pdf_photo(raw, target_width, target_height):
    """Center-crop a photo so it fills its PDF frame without distortion."""
    render_width = 720
    render_height = max(1, round(render_width * target_height / target_width))

    with Image.open(io.BytesIO(raw)) as source:
        source.load()
        oriented = ImageOps.exif_transpose(source)
        fitted = ImageOps.fit(
            oriented,
            (render_width, render_height),
            method=Image.Resampling.LANCZOS,
            centering=(0.5, 0.5),
        )

        if fitted.mode in ('RGBA', 'LA') or 'transparency' in fitted.info:
            rgba = fitted.convert('RGBA')
            flattened = Image.new('RGB', rgba.size, 'white')
            flattened.paste(rgba, mask=rgba.getchannel('A'))
            fitted = flattened
        elif fitted.mode != 'RGB':
            fitted = fitted.convert('RGB')

        output = io.BytesIO()
        fitted.save(output, format='JPEG', quality=88, optimize=True)
        output.seek(0)
        return output

def generate_pdf(d, output_path, cfg):
    st = make_styles(cfg)
    GREY_LINE = colors.HexColor('#CCCCCC')
    GREY_BG   = colors.HexColor('#F2F2F2')
    WHITE     = colors.white

    date  = d.get('date','');  day   = d.get('day_no','')
    proj  = d.get('project_no',''); loc = d.get('location','')
    cust  = d.get('customer','');   equip = d.get('equipment','-')
    prep  = d.get('prepared_by',''); chk = d.get('checked_by','')
    appr  = d.get('approved_by','')

    gpa_logo = resolve_logo_path(cfg.get('logo_gpa', ''), 'gpa')
    kn_logo = resolve_logo_path(cfg.get('logo_kn', ''), 'kn')
    gpa_is_bundled = bool(gpa_logo) and os.path.normcase(os.path.abspath(gpa_logo)) == \
        os.path.normcase(os.path.abspath(BUNDLED_LOGOS['gpa']))
    kn_is_bundled = bool(kn_logo) and os.path.normcase(os.path.abspath(kn_logo)) == \
        os.path.normcase(os.path.abspath(BUNDLED_LOGOS['kn']))
    use_gpa_header_logo = gpa_is_bundled and _is_drawable_logo(BUNDLED_HEADER_LOGOS['gpa'])
    use_kn_header_logo = kn_is_bundled and _is_drawable_logo(BUNDLED_HEADER_LOGOS['kn'])
    if use_gpa_header_logo:
        gpa_logo = BUNDLED_HEADER_LOGOS['gpa']
    if use_kn_header_logo:
        kn_logo = BUNDLED_HEADER_LOGOS['kn']

    _NC._meta = {
        'date': date, 'day': day, 'proj': proj, 'loc': loc, 'cust': cust,
        'theme':        cfg.get('theme', {}),
        'logo_gpa':     gpa_logo,
        'logo_kn':      kn_logo,
        'company_name': cfg.get('company_name',  'PT. GARUDA PRIMA AKSARA'),
        'project_title':d.get('project_title') or cfg.get('project_title','Electrical Installation & Construction'),
        'show_logo_gpa':cfg.get('show_logo_gpa', True),
        'show_logo_kn': cfg.get('show_logo_kn',  True),
        # The two bundled header marks use the same box and exact visual height.
        # Their widths remain proportional so neither brand mark is stretched.
        'logo_gpa_w':   43.5 if use_gpa_header_logo else cfg.get('logo_gpa_w', 28),
        'logo_gpa_h':   9.2 if use_gpa_header_logo else cfg.get('logo_gpa_h', 12),
        'logo_gpa_y_off':cfg.get('logo_gpa_y_off', 0),
        'logo_kn_w':    32 if use_kn_header_logo else cfg.get('logo_kn_w',  28),
        'logo_kn_h':    9.2 if use_kn_header_logo else cfg.get('logo_kn_h',  12),
        'logo_kn_y_off':cfg.get('logo_kn_y_off',  0),
        'logo_gpa_card':not use_gpa_header_logo,
        'logo_kn_card': not use_kn_header_logo,
    }

    story = []
    SECTION_GAP = 3.5 * mm

    def TC(value, centered=False):
        """Wrap user-entered table text instead of letting it cross a cell border."""
        safe_value = '' if value is None else _esc(value)
        return Paragraph(safe_value, st['tbl_c_s'] if centered else st['tbl_s'])

    def _bar(text, paragraph_style, background, height):
        """Build a full-width heading bar with vertically centred text."""
        bar = Table(
            [[Paragraph(text, paragraph_style)]],
            colWidths=[CW],
            rowHeights=[height],
        )
        bar.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), background),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('LEFTPADDING', (0, 0), (-1, -1), 3.5*mm),
            ('RIGHTPADDING', (0, 0), (-1, -1), 3*mm),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ]))
        # Tables below use ReportLab's centred overflow inside the padded page
        # frame. Match that placement so both outer edges line up exactly.
        bar.hAlign = 'CENTER'
        bar.keepWithNext = True
        return bar

    def SH(t):
        # A non-breaking gap keeps the section number readable without
        # changing the requested left alignment.
        safe = re.sub(r'^(\d+\.)\s*', r'\1&#160;&#160;', _esc(t))
        return [
            _bar(safe, st['sec_s'], st['PRI'], 6.5*mm),
            Spacer(1, 1.2*mm),
        ]

    def AH(t):
        safe = f'&#9632;&#160;&#160;{_esc(t)}'
        return [
            _bar(safe, st['area_s'], st['AREA'], 5.8*mm),
            Spacer(1, 1.2*mm),
        ]

    # S1 info
    story += SH('1.  REPORT INFORMATION')
    proj_title_val = d.get('project_title') or cfg.get('project_title', 'Electrical Installation &amp; Construction')
    info_rows = [
        ('Project No.',  _esc(proj)),
        ('Project Name', _esc(proj_title_val)),
        ('Customer',     _esc(cust)), ('Location', _esc(loc)), ('Equipment', _esc(equip)),
        ('Date', _esc(date)), ('Working Day', f'Day {_esc(day)}'),
        ('Working Hours','07:00 — 17:00 (Regular) | OT as noted'),
        ('Active Areas', '  '.join(_esc(a.get('id','')) for a in d.get('areas',[]))),
    ]
    itbl = Table([[Paragraph(k,S('Normal',fontSize=8,fontName='Helvetica-Bold')),Paragraph(v,st['body_s'])]
                  for k,v in info_rows], colWidths=[35*mm,CW-35*mm])
    itbl.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(0,-1),st['LB']),('GRID',(0,0),(-1,-1),0.4,GREY_LINE),
        ('FONTSIZE',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),2),
        ('BOTTOMPADDING',(0,0),(-1,-1),2),('LEFTPADDING',(0,0),(-1,-1),4),
        ('VALIGN',(0,0),(-1,-1),'MIDDLE'),('ROWBACKGROUNDS',(0,0),(-1,-1),[WHITE,GREY_BG])]))
    story += [itbl, Spacer(1,SECTION_GAP)]

    # S2 weather
    story += SH('2.  WEATHER REPORT')
    wx = d.get('weather',{})
    if wx:
        keys = list(wx.keys())
        vals = [TC(wx[k], centered=True) for k in keys]
        keys = [Paragraph(_esc(k), st['wx_h_s']) for k in keys]
        wt = Table([keys,vals], colWidths=[CW/len(keys)]*len(keys))
        wt.setStyle(TableStyle([
            ('BACKGROUND',(0,0),(-1,0),st['SEC']),('TEXTCOLOR',(0,0),(-1,0),WHITE),
            ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('FONTSIZE',(0,0),(-1,-1),7),
            ('GRID',(0,0),(-1,-1),0.4,GREY_LINE),('ALIGN',(0,0),(-1,-1),'CENTER'),
            ('VALIGN',(0,0),(-1,-1),'MIDDLE'),('TOPPADDING',(0,0),(-1,-1),2),
            ('BOTTOMPADDING',(0,0),(-1,-1),2)]))
        story.append(wt)
    story.append(Spacer(1,SECTION_GAP))

    # S3 indirect
    story += SH('3.  INDIRECT MANPOWER')
    ind = d.get('indirect_manpower',[])
    irows = [['No.','Name','Role / Position','Working Hours']]
    for i,p in enumerate(ind,1):
        irows.append([
            str(i), TC(p.get('name','')), TC(p.get('role','')),
            TC(p.get('hours',''), centered=True),
        ])
    itbl2 = Table(irows, colWidths=[9*mm,70*mm,55*mm,CW-134*mm])
    itbl2.setStyle(base_ts([
        ('ALIGN',(0,0),(0,-1),'CENTER'),('ALIGN',(3,0),(3,-1),'CENTER'),
        ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
    ]))
    story += [itbl2, Spacer(1,SECTION_GAP)]
    story.append(CondPageBreak(32*mm))

    section_number = 4
    show_overall_progress = _coerce_bool(d.get('show_overall_progress'), False)

    # Optional overall progress section. Missing flags default to disabled;
    # explicit true values from the current form still enable the section.
    if show_overall_progress:
        story += SH(f'{section_number}.  OVERALL PROGRESS')
        progress_rows = _normalise_overall_progress(d.get('overall_progress', []))
        if progress_rows:
            progress_h = S(
                'Normal', fontSize=5.1, leading=5.8, fontName='Helvetica-Bold',
                textColor=WHITE, alignment=1, splitLongWords=1,
            )
            progress_c = S(
                'Normal', fontSize=5.5, leading=6.4, alignment=1,
                splitLongWords=1,
            )
            progress_l = S(
                'Normal', fontSize=6, leading=7, splitLongWords=1,
            )
            progress_total = S(
                'Normal', fontSize=5.5, leading=6.4, fontName='Helvetica-Bold',
                alignment=1, splitLongWords=1,
            )

            def PH(value):
                return Paragraph(_esc(value), progress_h)

            def PC(value, left=False):
                return Paragraph(_esc(value), progress_l if left else progress_c)

            progress_data = [
                [
                    PH('No.'), PH('Description'), PH('Duration'), PH('Weight Factor'),
                    PH('Start'), PH('Finish'), PH('Cumulative Previous'), '',
                    PH('This Period'), '', PH('Cumulative Up to This Month'), '', '',
                ],
                ['', '', '', '', '', '', PH('Plan'), PH('Actual'), PH('Plan'),
                 PH('Actual'), PH('Plan'), PH('Actual'), PH('Deviation')],
            ]

            for index, row in enumerate(progress_rows, 1):
                deviation = row.get('deviation', '')
                if not deviation:
                    cumulative_plan = _progress_number(row.get('cumulative_to_date_plan'))
                    cumulative_actual = _progress_number(row.get('cumulative_to_date_actual'))
                    if cumulative_plan is not None and cumulative_actual is not None:
                        deviation = f'{cumulative_actual - cumulative_plan:g}'
                progress_data.append([
                    PC(index),
                    PC(row.get('description', ''), left=True),
                    PC(row.get('duration', '')),
                    PC(_progress_percent_text(row.get('weight_factor', ''))),
                    PC(row.get('start', '')),
                    PC(row.get('finish', '')),
                    *[
                        PC(_progress_percent_text(row.get(field, '')))
                        for field in OVERALL_PROGRESS_PERCENT_FIELDS
                    ],
                    PC(_progress_percent_text(deviation)),
                ])

            totals = _overall_progress_totals(progress_rows)
            progress_data.append([
                Paragraph('OVERALL PROGRESS', progress_total), '', '', '', '', '',
                *[
                    Paragraph(
                        '' if totals.get(field) is None else f"{totals[field]:.2f}%",
                        progress_total,
                    )
                    for field in OVERALL_PROGRESS_PERCENT_FIELDS
                ],
                Paragraph(
                    '' if totals.get('deviation') is None else f"{totals['deviation']:.2f}%",
                    progress_total,
                ),
            ])

            progress_widths = [
                7*mm, 53*mm, 10*mm, 11*mm, 14*mm, 14*mm,
                10*mm, 10*mm, 10*mm, 10*mm, 10*mm, 10*mm, 11*mm,
            ]
            progress_table = Table(
                progress_data,
                colWidths=progress_widths,
                repeatRows=2,
                splitByRow=1,
            )
            progress_table.setStyle(TableStyle([
                ('SPAN',(0,0),(0,1)),('SPAN',(1,0),(1,1)),('SPAN',(2,0),(2,1)),
                ('SPAN',(3,0),(3,1)),('SPAN',(4,0),(4,1)),('SPAN',(5,0),(5,1)),
                ('SPAN',(6,0),(7,0)),('SPAN',(8,0),(9,0)),('SPAN',(10,0),(12,0)),
                ('BACKGROUND',(0,0),(-1,1),st['PRI']),
                ('TEXTCOLOR',(0,0),(-1,1),WHITE),
                ('GRID',(0,0),(-1,-1),0.35,GREY_LINE),
                ('ROWBACKGROUNDS',(0,2),(-1,-2),[WHITE,GREY_BG]),
                ('BACKGROUND',(0,-1),(-1,-1),st['LB']),
                ('SPAN',(0,-1),(5,-1)),
                ('ALIGN',(0,0),(-1,-1),'CENTER'),
                ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
                ('TOPPADDING',(0,0),(-1,-1),1.2),
                ('BOTTOMPADDING',(0,0),(-1,-1),1.2),
                ('LEFTPADDING',(0,0),(-1,-1),1.2),
                ('RIGHTPADDING',(0,0),(-1,-1),1.2),
            ]))
            story.append(progress_table)
        else:
            story.append(Paragraph('No overall progress reported.', st['ital_s']))
        story.append(Spacer(1, SECTION_GAP))
        story.append(CondPageBreak(32*mm))
        section_number += 1

    # Daily activities and manpower by area
    story += SH(f'{section_number}.  DAILY ACTIVITIES & MANPOWER BY AREA')
    section_number += 1
    areas = d.get('areas', [])
    for area_index, area in enumerate(areas):
        aid = area.get('id','')
        blocks = list(AH(aid))

        def act_cell(lines, label):
            parts = [Paragraph(label, st['bold_s'])]
            for j,line in enumerate([l for l in lines if str(l).strip()],1):
                parts.append(Paragraph(f'{j}.  {_esc(line)}', st['sm_s']))
            if not parts[1:]:
                parts.append(Paragraph('—', st['ital_s']))
            return parts

        if area.get('activities_swapped'):
            cols = [act_cell(area.get('activities_tomorrow',[]),'Activity Tomorrow'),
                    act_cell(area.get('activities_today',[]),'Activity Today')]
        else:
            cols = [act_cell(area.get('activities_today',[]),'Activity Today'),
                    act_cell(area.get('activities_tomorrow',[]),'Activity Tomorrow')]
        at = Table([cols], colWidths=[CW*0.55,CW*0.45])
        at.setStyle(TableStyle([
            ('VALIGN',(0,0),(-1,-1),'TOP'),('BOX',(0,0),(-1,-1),0.5,GREY_LINE),
            ('INNERGRID',(0,0),(-1,-1),0.4,GREY_LINE),('BACKGROUND',(0,0),(-1,-1),GREY_BG),
            ('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3),
            ('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4)]))
        blocks += [at, Spacer(1,1*mm)]

        mp = area.get('manpower',[])
        if mp:
            blocks.append(Paragraph(f'Direct Manpower — {_esc(aid)}', st['sub_s']))
            mrows = [['No.','Name','Position','Task Today','Hours']]
            for j,p in enumerate(mp,1):
                mrows.append([
                    str(j), TC(p.get('name','')), TC(p.get('role','')),
                    TC(p.get('task','')), TC(p.get('hours',''), centered=True),
                ])
            mt = Table(mrows, colWidths=[8*mm,38*mm,30*mm,74*mm,CW-150*mm])
            mt.setStyle(base_ts([
                ('ALIGN',(0,0),(0,-1),'CENTER'),('ALIGN',(4,0),(4,-1),'CENTER'),
                ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
            ]))
            blocks += [mt, Spacer(1,1*mm)]

        # per-area indirect manpower
        area_ind = area.get('indirect_manpower', [])
        if area_ind:
            blocks.append(Paragraph(f'Indirect Manpower — {_esc(aid)}', st['sub_s']))
            irows_a = [['No.','Name','Role / Position','Working Hours']]
            for j, p in enumerate(area_ind, 1):
                irows_a.append([
                    str(j), TC(p.get('name','')), TC(p.get('role','')),
                    TC(p.get('hours',''), centered=True),
                ])
            it_a = Table(irows_a, colWidths=[9*mm, 70*mm, 55*mm, CW - 134*mm])
            it_a.setStyle(base_ts([('ALIGN',(0,0),(0,-1),'CENTER'),
                                    ('ALIGN',(3,0),(3,-1),'CENTER'),
                                    ('VALIGN',(0,0),(-1,-1),'MIDDLE')]))
            blocks += [it_a, Spacer(1,1*mm)]

        ct = area.get('constraints','').strip()
        rm = area.get('remarks','').strip()
        if ct or rm:
            cr = []
            if ct: cr.append([Paragraph('Constraints:',st['bold_s']),Paragraph(_esc(ct),st['body_s'])])
            if rm: cr.append([Paragraph('Remarks:',st['bold_s']),Paragraph(_esc(rm),st['body_s'])])
            crt = Table(cr,colWidths=[25*mm,CW-25*mm])
            crt.setStyle(TableStyle([
                ('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),0.3,GREY_LINE),
                ('BACKGROUND',(0,0),(0,-1),colors.HexColor('#FFF8E7')),
                ('TOPPADDING',(0,0),(-1,-1),2),('BOTTOMPADDING',(0,0),(-1,-1),2),
                ('LEFTPADDING',(0,0),(-1,-1),4)]))
            blocks += [crt, Spacer(1,1*mm)]

        # The area-to-area gap is added below so every area ends with the same
        # amount of whitespace, regardless of which optional tables it has.
        if blocks and isinstance(blocks[-1], Spacer):
            blocks.pop()
        story.append(KeepTogether(blocks[:4]))
        for b in blocks[4:]: story.append(b)
        if area_index < len(areas) - 1:
            story.append(Spacer(1,2.5*mm))

    story.append(Spacer(1,SECTION_GAP))
    # The compact constraints/remarks pair needs about 25 mm. Keeping this
    # threshold precise prevents a taller page header from wasting the rest
    # of page one and pushing the final photo row onto a third page.
    story.append(CondPageBreak(25*mm))

    # Constraints and issues
    story += SH(f'{section_number}.  CONSTRAINTS & ISSUES')
    section_number += 1
    any_c = any(a.get('constraints','').strip() for a in d.get('areas',[]))
    if any_c:
        gcr = [['Area','Constraint / Issue']]
        for a in d.get('areas',[]):
            c = a.get('constraints','').strip()
            if c: gcr.append([TC(a.get('id','')),TC(c)])
        gct = Table(gcr,colWidths=[20*mm,CW-20*mm])
        gct.setStyle(base_ts([('VALIGN',(0,0),(-1,-1),'TOP')]))
        story.append(gct)
    else:
        story.append(Paragraph('No constraints reported.',st['ital_s']))
    story.append(Spacer(1,SECTION_GAP))

    # Remarks
    story += SH(f'{section_number}.  REMARKS')
    section_number += 1
    gr = d.get('global_remarks','').strip()
    story += [Paragraph(_esc(gr) if gr else '—',st['body_s']),Spacer(1,SECTION_GAP)]
    story.append(CondPageBreak(42*mm))

    # Sign-off (dynamic columns)
    story += SH(f'{section_number}.  SIGN-OFF')
    section_number += 1
    sign_offs = d.get('sign_offs', [])
    if not sign_offs:
        sign_offs = [
            {'label':'Prepared By',       'name': prep, 'role':'HSE / Administration', 'sig':''},
            {'label':'Checked By',        'name': chk,  'role':'Project Control',       'sig':''},
            {'label':'Approved By',       'name': appr, 'role':'Project Manager',       'sig':''},
            {'label':'KN Representative', 'name': '',   'role':'PT. Kertas Nusantara',  'sig':''},
        ]
    n_cols = len(sign_offs)
    col_w  = CW / max(n_cols, 1)
    def _sig_cell(sd):
        if sd and ',' in sd:
            try:
                ib = io.BytesIO(base64.b64decode(sd.split(',')[1]))
                return RLImage(ib, width=min(col_w - 10*mm, 42*mm), height=20*mm)
            except: pass
        return Spacer(1, 20*mm)
    _sh = ParagraphStyle('_sh', fontName='Helvetica-Bold', fontSize=8, textColor=colors.white, alignment=1)
    _sr = ParagraphStyle('_sr', fontName='Helvetica', fontSize=7, textColor=colors.grey, alignment=1)
    so = Table([
        [Paragraph(_esc(s.get('label','')), _sh)        for s in sign_offs],
        [_sig_cell(s.get('sig',''))                      for s in sign_offs],
        [Paragraph('_' * 22, st['body_s'])               for s in sign_offs],
        [Paragraph(_esc(s.get('name','')), st['bold_s']) for s in sign_offs],
        [Paragraph(_esc(s.get('role','')), _sr)          for s in sign_offs],
    ], colWidths=[col_w] * n_cols)
    so.setStyle(TableStyle([
        ('BACKGROUND',(0,0),(-1,0),st['PRI']),('TEXTCOLOR',(0,0),(-1,0),WHITE),
        ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('FONTSIZE',(0,0),(-1,-1),8.5),
        ('ALIGN',(0,0),(-1,-1),'CENTER'),('VALIGN',(0,0),(-1,-1),'MIDDLE'),
        ('GRID',(0,0),(-1,-1),0.5,GREY_LINE),
        ('TOPPADDING',(0,0),(-1,-1),2),('BOTTOMPADDING',(0,0),(-1,-1),2),
        ('FONTNAME',(0,3),(-1,3),'Helvetica-Bold'),
        ('FONTSIZE',(0,4),(-1,4),7.5),('TEXTCOLOR',(0,4),(-1,4),colors.grey)]))
    story += [so, Spacer(1,SECTION_GAP)]
    PER_ROW = 3
    photo_sections = _group_report_photos(areas, PER_ROW)
    story.append(CondPageBreak(65*mm if photo_sections else 10*mm))

    # Photo documentation. Each area has its own heading and grid so photos from different
    # work areas are never presented as one continuous group.
    story += SH(f'{section_number}.  PHOTO DOCUMENTATION')
    photo_documentation_title = _bounded_pdf_text(
        d.get('photo_documentation_title', ''),
        PDF_PHOTO_AREA_MAX_CHARS,
    )
    if photo_documentation_title:
        story.append(Paragraph(_esc(photo_documentation_title), st['sub_s']))
        story.append(Spacer(1, 0.5*mm))
    BOX_W=(CW-4*mm)/PER_ROW; BOX_H=52*mm
    PHOTO_INSET = 0
    text_width = BOX_W - 10*mm
    for raw_area_id, photo_groups in photo_sections:
        aid = _bounded_pdf_text(raw_area_id, PDF_PHOTO_AREA_MAX_CHARS)
        story.append(CondPageBreak(65*mm))
        story.append(Paragraph(_esc(aid), st['sub_s']))
        story.append(Spacer(1, 0.5*mm))
        rows = []

        for photo_group in photo_groups:
            group = []
            for _, photo in photo_group:
                img_data = photo.get('img_data','')
                desc = _bounded_pdf_text(photo.get('desc', ''), PDF_PHOTO_CAPTION_MAX_CHARS)
                if img_data and ',' in img_data:
                    try:
                        b64 = img_data.split(',')[1]
                        raw_photo = base64.b64decode(b64)
                        photo_w = BOX_W - 4*mm - 2*PHOTO_INSET
                        photo_h = BOX_H - 2*PHOTO_INSET
                        img_bytes = _prepare_pdf_photo(raw_photo, photo_w, photo_h)
                        rl_img = RLImage(img_bytes, width=photo_w, height=photo_h)
                        photo_cell = rl_img
                    except:
                        photo_cell = ''
                else:
                    photo_cell = ''
                card_title = photo_documentation_title or aid
                title_para = Paragraph(f'<b>{_esc(card_title)}</b>', st['sm_s'])
                caption_para = Paragraph(_esc(desc), st['ital_s'])
                group.append((title_para, caption_para, photo_cell))

            # Let long area names and captions wrap. Every card in one area row
            # uses the tallest text block so all photo borders remain aligned.
            title_h = max(
                5*mm,
                max(title.wrap(text_width, 100*mm)[1] for title, _, _ in group) + 2*mm,
            )
            caption_h = max(
                6*mm,
                max(caption.wrap(text_width, 100*mm)[1] for _, caption, _ in group) + 2*mm,
            )

            row = []
            for title_para, caption_para, photo_cell in group:
                inner = Table(
                    [[title_para], [caption_para], [photo_cell]],
                    colWidths=[BOX_W-4*mm],
                    rowHeights=[title_h, caption_h, BOX_H],
                )
                inner.setStyle(TableStyle([
                    ('BOX',(0,0),(-1,-1),0.8,st['PRI']),
                    ('LINEBELOW',(0,0),(0,0),0.5,GREY_LINE),
                    ('LINEBELOW',(0,1),(0,1),0.5,GREY_LINE),
                    ('BACKGROUND',(0,0),(0,0),st['LB']),
                    ('BACKGROUND',(0,1),(0,1),GREY_BG),
                    # The image reaches the inside edge of the blue frame. A
                    # matching background prevents white hairlines from PDF
                    # sub-pixel rounding, while the vector border stays crisp.
                    ('BACKGROUND',(0,2),(0,2),st['PRI'] if photo_cell else WHITE),
                    ('TOPPADDING',(0,0),(-1,1),2),('BOTTOMPADDING',(0,0),(-1,1),2),
                    ('LEFTPADDING',(0,0),(-1,1),3),('RIGHTPADDING',(0,0),(-1,1),3),
                    ('TOPPADDING',(0,2),(-1,2),PHOTO_INSET),
                    ('BOTTOMPADDING',(0,2),(-1,2),PHOTO_INSET),
                    ('LEFTPADDING',(0,2),(-1,2),PHOTO_INSET),
                    ('RIGHTPADDING',(0,2),(-1,2),PHOTO_INSET),
                    ('VALIGN',(0,0),(-1,1),'MIDDLE'),
                ]))
                row.append(inner)

            while len(row) < PER_ROW:
                row.append('')
            rows.append(row)

        # Each area's rows get their own table, so the next area always starts
        # as a visibly separate photo group.
        pg=Table(rows,colWidths=[BOX_W]*PER_ROW)
        pg.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),
            ('LEFTPADDING',(0,0),(-1,-1),2),('RIGHTPADDING',(0,0),(-1,-1),2),
            ('TOPPADDING',(0,0),(-1,-1),2),('BOTTOMPADDING',(0,0),(-1,-1),4)]))
        story += [pg, Spacer(1,2*mm)]

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf,pagesize=A4,topMargin=40.5*mm,bottomMargin=11*mm,leftMargin=M,rightMargin=M)
    doc.build(story, canvasmaker=_NC)
    if output_path:
        os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else '.', exist_ok=True)
        with open(output_path,'wb') as f: f.write(buf.getvalue())
    buf.seek(0)
    return buf

# ── Data ──────────────────────────────────────────────────────────────────────
MANPOWER_DB = [
    # Indirect
    {"name":"Faiz Satria",                   "role":"Project Control",          "type":"indirect"},
    {"name":"Sodiq",                          "role":"Document Control",         "type":"indirect"},
    {"name":"Zulfikar",                       "role":"GA",                       "type":"indirect"},
    {"name":"Nafiz",                          "role":"Administration",            "type":"indirect"},
    {"name":"Trivena Kasih Kristiani",        "role":"Admin Site",               "type":"indirect"},
    {"name":"Phasa Amalia Arzetti Putri",     "role":"Admin",                    "type":"indirect"},
    # Direct
    {"name":"Asril Naimy",                    "role":"Project Manager",          "type":"direct"},
    {"name":"Pipin Arifin",                   "role":"Site Manager",             "type":"direct"},
    {"name":"Eddy Suyardi",                   "role":"Supervisor",               "type":"direct"},
    {"name":"Siswono",                        "role":"Supervisor",               "type":"direct"},
    {"name":"Rusdyanto",                      "role":"Supervisor",               "type":"direct"},
    {"name":"Hargo Wahono Edy.S",             "role":"Supervisor",               "type":"direct"},
    {"name":"Vemy Amelia",                    "role":"HSE",                      "type":"direct"},
    {"name":"Deby Sandra",                    "role":"HSE",                      "type":"direct"},
    {"name":"Hamkah",                         "role":"Foreman",                  "type":"direct"},
    {"name":"Maryono",                        "role":"Foreman",                  "type":"direct"},
    {"name":"Marjulisman",                    "role":"Foreman",                  "type":"direct"},
    {"name":"Riza Amri",                      "role":"Foreman",                  "type":"direct"},
    {"name":"Arief Maulana",                  "role":"Foreman",                  "type":"direct"},
    {"name":"Sanudin",                        "role":"Foreman",                  "type":"direct"},
    {"name":"Surya Purwanto",                 "role":"Foreman",                  "type":"direct"},
    {"name":"Rudy Safutra",                   "role":"Fitter (Teknisi)",         "type":"direct"},
    {"name":"Fadillah",                       "role":"Teknisi",                  "type":"direct"},
    {"name":"Feriyanto",                      "role":"Teknisi",                  "type":"direct"},
    {"name":"Syafril",                        "role":"Teknisi",                  "type":"direct"},
    {"name":"Firman",                         "role":"Teknisi",                  "type":"direct"},
    {"name":"Ali Iskandar",                   "role":"Teknisi",                  "type":"direct"},
    {"name":"Nova",                           "role":"Teknisi",                  "type":"direct"},
    {"name":"Warnoto",                        "role":"Teknisi",                  "type":"direct"},
    {"name":"Yanto",                          "role":"Teknisi",                  "type":"direct"},
    {"name":"Yusuf",                          "role":"Teknisi",                  "type":"direct"},
    {"name":"Tri Ageng",                      "role":"Teknisi",                  "type":"direct"},
    {"name":"Timbul",                         "role":"Teknisi",                  "type":"direct"},
    {"name":"Ahmad Haris",                    "role":"Teknisi",                  "type":"direct"},
    {"name":"Trio Kurniawan",                 "role":"Teknisi",                  "type":"direct"},
    {"name":"Yasmin",                         "role":"Teknisi",                  "type":"direct"},
    {"name":"Fuji",                           "role":"Teknisi",                  "type":"direct"},
    {"name":"Kardi",                          "role":"Teknisi",                  "type":"direct"},
    {"name":"Firdaus",                        "role":"Teknisi",                  "type":"direct"},
    {"name":"Ikbal Mariado",                  "role":"Teknisi",                  "type":"direct"},
    {"name":"Andi Iskandar Muda",             "role":"Teknisi",                  "type":"direct"},
    {"name":"Irwan Suryanto",                 "role":"Teknisi",                  "type":"direct"},
    {"name":"Riswan",                         "role":"Teknisi",                  "type":"direct"},
    {"name":"Yusril Mahendra",                "role":"Teknisi",                  "type":"direct"},
    {"name":"Ramang",                         "role":"Teknisi",                  "type":"direct"},
    {"name":"Afrizal",                        "role":"Teknisi",                  "type":"direct"},
    {"name":"Iwan S",                         "role":"Teknisi",                  "type":"direct"},
    {"name":"Agus Sulistiyo",                 "role":"Teknisi",                  "type":"direct"},
    {"name":"Suhadi",                         "role":"Teknisi",                  "type":"direct"},
    {"name":"Bambang Setyawan",               "role":"Teknisi",                  "type":"direct"},
    {"name":"Yudi Ivanto",                    "role":"Teknisi",                  "type":"direct"},
    {"name":"Rafi Ahmad Pradana",             "role":"Teknisi",                  "type":"direct"},
    {"name":"Decmanto Canda Lalong",          "role":"Teknisi",                  "type":"direct"},
    {"name":"Edy Tavip",                      "role":"Teknisi",                  "type":"direct"},
    {"name":"Alfian",                         "role":"Teknisi",                  "type":"direct"},
    {"name":"Agustinus Allodatu",             "role":"Teknisi",                  "type":"direct"},
    {"name":"Syaharuddin",                    "role":"Teknisi",                  "type":"direct"},
    {"name":"Solihin",                        "role":"Teknisi",                  "type":"direct"},
    {"name":"Ignatius Alfatendo Putra Fau",   "role":"Teknisi",                  "type":"direct"},
    {"name":"Fasha Muhamad Rizki",            "role":"Teknisi",                  "type":"direct"},
    {"name":"Roby",                           "role":"Teknisi",                  "type":"direct"},
    {"name":"Muhamad Dede Saputra",           "role":"Teknisi",                  "type":"direct"},
    {"name":"Hilman F",                       "role":"Teknisi",                  "type":"direct"},
    {"name":"Yohanis Tandililing",            "role":"Welder (Teknisi)",         "type":"direct"},
    {"name":"Dimas",                          "role":"Welder (Teknisi)",         "type":"direct"},
    {"name":"Aris Seno",                      "role":"Teknisi",                  "type":"direct"},
    {"name":"Wiranto Mangin",                 "role":"Welder (Teknisi)",         "type":"direct"},
    {"name":"Gregorius Rovino Batu",          "role":"Helper",                   "type":"direct"},
    {"name":"Yarius T",                       "role":"Helper",                   "type":"direct"},
    {"name":"Kelvinus Siregar",               "role":"Helper",                   "type":"direct"},
    {"name":"Paulus Raba'",                   "role":"Helper",                   "type":"direct"},
    {"name":"Yonatan Devi Riantori",          "role":"Helper",                   "type":"direct"},
    {"name":"Abraham Londong",                "role":"Helper",                   "type":"direct"},
    {"name":"Gedofrianto Gunawan",            "role":"Helper",                   "type":"direct"},
    {"name":"Agustinus Dion Eba",             "role":"Helper",                   "type":"direct"},
    {"name":"Zainuddin",                      "role":"Helper",                   "type":"direct"},
    {"name":"Yosep Riki Hermanto",            "role":"Helper",                   "type":"direct"},
    {"name":"Rio Saputra",                    "role":"Helper",                   "type":"direct"},
    {"name":"Zeth Pabontong",                 "role":"Helper",                   "type":"direct"},
    {"name":"Versianus Bahanu",               "role":"Helper",                   "type":"direct"},
    {"name":"Yusuf Mappa",                    "role":"Helper",                   "type":"direct"},
    {"name":"Ricky Irfan",                    "role":"Helper",                   "type":"direct"},
    {"name":"Resa Saputra",                   "role":"Helper",                   "type":"direct"},
    {"name":"Arill",                          "role":"Helper",                   "type":"direct"},
    {"name":"Hardi",                          "role":"Teknisi (Scaffolding)",    "type":"direct"},
    {"name":"Novrianto Rangan",               "role":"Helper",                   "type":"direct"},
    {"name":"Rafly Hergenveral Sampe",        "role":"Helper",                   "type":"direct"},
    {"name":"Rahmat",                         "role":"Helper",                   "type":"direct"},
    {"name":"Jusli",                          "role":"Helper",                   "type":"direct"},
    {"name":"Tandi",                          "role":"Helper",                   "type":"direct"},
    {"name":"Yahdillah",                      "role":"Helper",                   "type":"direct"},
    {"name":"William",                        "role":"Helper",                   "type":"direct"},
    {"name":"Registor Enga",                  "role":"Helper",                   "type":"direct"},
    {"name":"Jemmy Juwardi As",               "role":"Helper",                   "type":"direct"},
    {"name":"Bayu Setiawan",                  "role":"Helper",                   "type":"direct"},
    {"name":"Kevin Garanta",                  "role":"Helper",                   "type":"direct"},
    {"name":"Yusrianto",                      "role":"Helper",                   "type":"direct"},
    {"name":"Randy",                          "role":"Helper",                   "type":"direct"},
    {"name":"Sandi Anugrah",                  "role":"Helper",                   "type":"direct"},
    {"name":"Samsuddin",                      "role":"Helper",                   "type":"direct"},
    {"name":"M. Jepriansyah",                 "role":"Tool Keeper",              "type":"direct"},
    {"name":"Aas Dani Miharja",               "role":"Teknisi",                  "type":"direct"},
    {"name":"RD. Ilham Abillah Pangestu",     "role":"Teknisi",                  "type":"direct"},
    {"name":"Romei Hendianto",                "role":"Teknisi",                  "type":"direct"},
    {"name":"Edwan Sukmayana",                "role":"Teknisi",                  "type":"direct"},
    {"name":"Irfan Akbar",                    "role":"Teknisi",                  "type":"direct"},
    {"name":"Muhamad Said",                   "role":"Teknisi",                  "type":"direct"},
    {"name":"Risman",                         "role":"Helper",                   "type":"direct"},
    {"name":"Bartholomeus",                   "role":"Helper",                   "type":"direct"},
]

AREA_LIST = ["MA-14","MA-23","MA-24","MA-26","MA-39","MA-40","MA-41","MA-42",
             "MA-59","MA-73","MA-77","MA-81","MA-85"]

DEFAULT_PROJECTS = [
    {
        "title": "Electrical Construction and Installation - Manpower Supply",
        "project_no": "002/KN-GPA/EPC-2K-P2/XI/2025",
    },
    {
        "title": "Repair & Services Control Valve & ON OFF Valve",
        "project_no": "P01.0825.J075",
    },
    {
        "title": "PROJECT REVAMPING PT KERTAS NUSANTARA - REACTIVATION FOR TURBINES AND GENERATORS",
        "project_no": "001/KN-GPA/EPC-2F-P2/IV/2025",
    },
]

DEFAULT_CONFIG = {
    "company_name": "PT. GARUDA PRIMA AKSARA",
    "customer": "PT. KERTAS NUSANTARA",
    "project_no": "PC-26-0004-KN-GPA-029-DAR",
    "location": "Berau, East Kalimantan",
    "equipment": "-",
    "prepared_by": "Phasa Amalia Arzetti Putri",
    "checked_by": "Faiz M. Satria N",
    "approved_by": "Asril Naimy",
    "logo_gpa": "",
    "logo_kn":  "",
    "project_title": "Electrical Installation & Construction",
    "projects": DEFAULT_PROJECTS,
    "show_logo_gpa":   True, "show_logo_kn":   True,
    "logo_gpa_w": 28,  "logo_gpa_h": 12,  "logo_gpa_y_off": 0,
    "logo_kn_w":  28,  "logo_kn_h":  12,  "logo_kn_y_off":  0,
    "theme": {
        "primary":   "#003366",
        "secondary": "#005B99",
        "accent":    "#C89010",
        "area_hdr":  "#1A5276",
        "light_bg":  "#D6E8F7",
        "navbar_bg": "#003366"
    },
    "areas": AREA_LIST,
    "hours_options": [
        "07:00 - 17:00","07:00 - 18:00","07:00 - 19:00",
        "07:00 - 20:00","07:00 - 21:00","07:00 - 22:00"
    ],
    "project_start_date": "",
    "letter_seq_no": 1,
    "pm_name": "Asril Naimy",
    "pm_title": "Project Manager",
    "site_name": "Mangkajang",
    "letter_kn_attn": "",
    "manpower_db": MANPOWER_DB,
    "ai_api_key": "",
}

# ── Flask ─────────────────────────────────────────────────────────────────────
import hashlib, functools

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'gpa-daily-report-s3cr3t-2026')

SCRIPT_DIR        = os.path.dirname(os.path.abspath(__file__))
# Keep mutable data outside the deployed source when DATA_DIR is configured.
# Railway should mount a persistent Volume at this path (for example /data).
DATA_DIR          = os.path.abspath(os.environ.get('DATA_DIR', SCRIPT_DIR))
CONFIG_FILE       = os.path.join(DATA_DIR, 'app_config.json')
LOGOS_DIR         = os.path.join(DATA_DIR, 'logos')
USERS_FILE        = os.path.join(DATA_DIR, 'users.json')
USERS_DIR         = os.path.join(DATA_DIR, 'users')
ACTIVITY_LOG_FILE = os.path.join(DATA_DIR, 'activity_log.json')
os.makedirs(DATA_DIR, exist_ok=True)
os.makedirs(LOGOS_DIR, exist_ok=True)
os.makedirs(USERS_DIR, exist_ok=True)
os.makedirs(os.path.join(SCRIPT_DIR, 'templates'), exist_ok=True)

# ── User / auth helpers ───────────────────────────────────────────────────────
def hash_pin(pin):
    return hashlib.sha256(str(pin).encode()).hexdigest()

def load_users():
    if os.path.exists(USERS_FILE):
        with open(USERS_FILE) as f:
            return json.load(f)
    initial_admin_pin = os.environ.get('ADMIN_PIN', '0000')
    users = {'admin': {'pin_hash': hash_pin(initial_admin_pin), 'is_admin': True,
                       'created_at': datetime.now().strftime('%Y-%m-%d')}}
    _save_users(users)
    return users

def _save_users(users):
    with open(USERS_FILE, 'w') as f:
        json.dump(users, f, indent=2)

def get_user_dir(username):
    d = os.path.join(USERS_DIR, username)
    os.makedirs(os.path.join(d, 'reports'), exist_ok=True)
    return d

def get_temp_photos_dir(username):
    d = os.path.join(get_user_dir(username), 'temp_photos')
    os.makedirs(d, exist_ok=True)
    return d

def get_draft_file(username):
    return os.path.join(get_user_dir(username), 'draft.json')

def get_reports_dir(username):
    return os.path.join(get_user_dir(username), 'reports')

def get_reports_index(username):
    idx = os.path.join(get_reports_dir(username), 'index.json')
    if os.path.exists(idx):
        try:
            with open(idx, encoding='utf-8') as f:
                data = json.load(f)
                return data if isinstance(data, list) else []
        except: pass
    return []

def append_report_index(username, entry):
    idx = get_reports_index(username)
    idx.insert(0, entry)
    reports_dir = get_reports_dir(username)
    os.makedirs(reports_dir, exist_ok=True)
    index_path = os.path.join(reports_dir, 'index.json')
    temp_path = f'{index_path}.{uuid.uuid4().hex}.tmp'
    try:
        with open(temp_path, 'w', encoding='utf-8') as f:
            json.dump(idx, f, ensure_ascii=False, indent=2)
        os.replace(temp_path, index_path)
    finally:
        if os.path.exists(temp_path):
            try: os.remove(temp_path)
            except OSError: pass

def _safe_report_filename_part(value, fallback):
    value = str(value or fallback)
    value = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', '-', value)
    return value.strip(' ._-') or fallback

def archive_generated_report(username, filename, pdf_bytes, entry):
    """Atomically save a PDF copy, then add it to the My Reports index."""
    reports_dir = get_reports_dir(username)
    os.makedirs(reports_dir, exist_ok=True)
    report_path = os.path.join(reports_dir, filename)
    temp_path = f'{report_path}.{uuid.uuid4().hex}.tmp'
    try:
        with open(temp_path, 'wb') as f:
            f.write(pdf_bytes)
        os.replace(temp_path, report_path)
    finally:
        if os.path.exists(temp_path):
            try: os.remove(temp_path)
            except OSError: pass

    entry = dict(entry)
    entry['size_kb'] = round(len(pdf_bytes) / 1024, 1)
    append_report_index(username, entry)

def login_required(f):
    @functools.wraps(f)
    def decorated(*args, **kwargs):
        if 'username' not in session:
            return redirect(url_for('login_page'))
        return f(*args, **kwargs)
    return decorated

def admin_required(f):
    @functools.wraps(f)
    def decorated(*args, **kwargs):
        if 'username' not in session:
            return redirect(url_for('login_page'))
        users = load_users()
        if not users.get(session['username'], {}).get('is_admin'):
            return redirect(url_for('index'))
        return f(*args, **kwargs)
    return decorated

def normalize_projects(value, strict=False):
    """Return safe title/number pairs while allowing a title to have many numbers."""
    if not isinstance(value, list):
        if strict:
            raise ValueError('Projects must be a list.')
        return copy.deepcopy(DEFAULT_PROJECTS)
    if len(value) > 100:
        if strict:
            raise ValueError('A maximum of 100 projects is allowed.')
        value = value[:100]

    normalized = []
    seen = set()
    for entry in value:
        if not isinstance(entry, dict):
            if strict:
                raise ValueError('Each project must contain a title and project number.')
            continue
        title = entry.get('title', '')
        project_no = entry.get('project_no', entry.get('number', ''))
        if not isinstance(title, str) or not isinstance(project_no, str):
            if strict:
                raise ValueError('Project title and number must be text.')
            continue
        title = title.strip()
        project_no = project_no.strip()
        if not title or not project_no:
            if strict:
                raise ValueError('Project title and number cannot be empty.')
            continue
        if len(title) > 300 or len(project_no) > 150:
            if strict:
                raise ValueError('Project title or number is too long.')
            continue
        pair_key = (title.casefold(), project_no.casefold())
        if pair_key in seen:
            if strict:
                raise ValueError('Duplicate project title and number pair.')
            continue
        seen.add(pair_key)
        normalized.append({'title': title, 'project_no': project_no})
    return normalized

def load_config():
    defaults = copy.deepcopy(DEFAULT_CONFIG)
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, encoding='utf-8') as f:
                c = json.load(f)
            if not isinstance(c, dict):
                raise ValueError('Config root must be an object.')
            for k,v in defaults.items():
                if k not in c:
                    c[k] = copy.deepcopy(v)
                elif isinstance(v, dict):
                    if not isinstance(c[k], dict):
                        c[k] = copy.deepcopy(v)
                    else:
                        for kk,vv in v.items():
                            if kk not in c[k]: c[k][kk]=copy.deepcopy(vv)
            c['projects'] = normalize_projects(c.get('projects', defaults['projects']))
            return c
        except: pass
    return defaults

def save_config(c):
    temp_path = f'{CONFIG_FILE}.{uuid.uuid4().hex}.tmp'
    try:
        with open(temp_path, 'w', encoding='utf-8') as f:
            json.dump(c, f, ensure_ascii=False, indent=2)
        os.replace(temp_path, CONFIG_FILE)
    finally:
        if os.path.exists(temp_path):
            try: os.remove(temp_path)
            except OSError: pass

# ── Login / logout routes ────────────────────────────────────────────────────
@app.route('/login', methods=['GET','POST'])
def login_page():
    if request.method == 'POST':
        username = (request.form.get('username') or '').strip().lower()
        pin      = (request.form.get('pin') or '').strip()
        users    = load_users()
        u        = users.get(username)
        if u and u.get('pin_hash') == hash_pin(pin):
            session['username'] = username
            session['is_admin'] = u.get('is_admin', False)
            return redirect(url_for('index'))
        return render_template('login.html', error='Incorrect username or PIN.')
    if 'username' in session:
        return redirect(url_for('index'))
    return render_template('login.html', error=None)

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login_page'))

# ── Admin routes ─────────────────────────────────────────────────────────────
@app.route('/admin')
@admin_required
def admin_page():
    users = load_users()
    user_data = []
    for uname, info in users.items():
        reports = get_reports_index(uname)
        user_data.append({
            'username':   uname,
            'is_admin':   info.get('is_admin', False),
            'created_at': info.get('created_at', ''),
            'report_count': len(reports),
        })
    return render_template('admin.html',
        users=user_data,
        current_user=session['username'])

@app.route('/admin/add_user', methods=['POST'])
@admin_required
def admin_add_user():
    data = request.json
    username = (data.get('username') or '').strip().lower()
    pin      = str(data.get('pin') or '0000').strip()
    is_admin = bool(data.get('is_admin', False))
    if not username or len(username) < 2:
        return jsonify({'error': 'Username must be at least 2 characters.'}), 400
    users = load_users()
    if username in users:
        return jsonify({'error': f'User "{username}" already exists.'}), 400
    users[username] = {
        'pin_hash':   hash_pin(pin),
        'is_admin':   is_admin,
        'created_at': datetime.now().strftime('%Y-%m-%d'),
    }
    _save_users(users)
    get_user_dir(username)
    return jsonify({'ok': True})

@app.route('/admin/remove_user', methods=['POST'])
@admin_required
def admin_remove_user():
    username = (request.json.get('username') or '').strip().lower()
    if username == session['username']:
        return jsonify({'error': 'Cannot remove yourself.'}), 400
    users = load_users()
    if username not in users:
        return jsonify({'error': 'User not found.'}), 404
    del users[username]
    _save_users(users)
    return jsonify({'ok': True})

@app.route('/admin/reset_pin', methods=['POST'])
@admin_required
def admin_reset_pin():
    data = request.json
    username = (data.get('username') or '').strip().lower()
    new_pin  = str(data.get('pin') or '0000').strip()
    users = load_users()
    if username not in users:
        return jsonify({'error': 'User not found.'}), 404
    users[username]['pin_hash'] = hash_pin(new_pin)
    _save_users(users)
    return jsonify({'ok': True})

@app.route('/admin/user_reports/<username>')
@admin_required
def admin_user_reports(username):
    reports = get_reports_index(username)
    return jsonify({'username': username, 'reports': reports})

@app.route('/admin/download/<username>/<path:filename>')
@admin_required
def admin_download_report(username, filename):
    rdir = get_reports_dir(username)
    fpath = os.path.join(rdir, filename)
    if not os.path.isfile(fpath):
        return 'Not found', 404
    return send_file(fpath, as_attachment=True, download_name=filename)

@app.route('/admin/delete_report', methods=['POST'])
@admin_required
def admin_delete_report():
    data = request.json
    username = (data.get('username') or '').strip()
    filename = (data.get('filename') or '').strip()
    if not username or not filename:
        return jsonify({'error': 'Missing params'}), 400
    fpath = os.path.join(get_reports_dir(username), filename)
    if os.path.isfile(fpath):
        os.remove(fpath)
    idx = [r for r in get_reports_index(username) if r.get('filename') != filename]
    with open(os.path.join(get_reports_dir(username), 'index.json'), 'w') as f:
        json.dump(idx, f, indent=2)
    return jsonify({'ok': True})

# ── User report history ──────────────────────────────────────────────────────
@app.route('/my_reports')
@login_required
def my_reports():
    username = session['username']
    cfg = load_config()
    reports  = get_reports_index(username)
    return render_template('reports.html',
        reports=reports,
        username=username,
        project_start_date=cfg.get('project_start_date',''),
        is_admin=session.get('is_admin', False))

@app.route('/reports/download/<path:filename>')
@login_required
def download_report(filename):
    username = session['username']
    fpath = os.path.join(get_reports_dir(username), filename)
    if not os.path.isfile(fpath):
        return 'Not found', 404
    return send_file(fpath, as_attachment=True, download_name=filename)

@app.route('/reports/delete', methods=['POST'])
@login_required
def delete_report():
    username = session['username']
    filename = (request.json.get('filename') or '').strip()
    fpath = os.path.join(get_reports_dir(username), filename)
    if os.path.isfile(fpath):
        os.remove(fpath)
    idx = [r for r in get_reports_index(username) if r.get('filename') != filename]
    with open(os.path.join(get_reports_dir(username), 'index.json'), 'w') as f:
        json.dump(idx, f, indent=2)
    return jsonify({'ok': True})

@app.route('/reports/check_date')
@login_required
def check_report_date():
    username = session['username']
    date = request.args.get('date', '').strip()
    if not date:
        return jsonify({'exists': False})
    for entry in get_reports_index(username):
        if entry.get('date') == date:
            return jsonify({'exists': True, 'filename': entry.get('filename', date)})
    return jsonify({'exists': False})

@app.route('/health')
def health():
    import sys, importlib
    checks = {}
    for mod in ['flask', 'reportlab', 'PIL', 'json', 'os', 'io', 'base64']:
        try:
            importlib.import_module(mod)
            checks[mod] = 'OK'
        except ImportError as e:
            checks[mod] = f'MISSING: {e}'
    checks['python'] = sys.version
    checks['templates_dir'] = os.path.exists(
        os.path.join(os.path.dirname(__file__), 'templates'))
    checks['index_html'] = os.path.exists(
        os.path.join(os.path.dirname(__file__), 'templates', 'index.html'))
    return jsonify(checks)

@app.route('/')
@login_required
def index():
    username = session['username']
    cfg = load_config()
    draft = {}
    draft_file = get_draft_file(username)
    if os.path.exists(draft_file):
        try:
            with open(draft_file) as f: draft = json.load(f)
        except: pass
    return render_template('index.html',
        manpower_db=cfg.get('manpower_db', MANPOWER_DB),
        area_list=cfg.get('areas', AREA_LIST),
        hours_options=cfg.get('hours_options', DEFAULT_CONFIG['hours_options']),
        app_config=cfg,
        initial_data=draft or None,
        username=username,
        is_admin=session.get('is_admin', False),
    )

# ── Photo server helpers ──────────────────────────────────────────────────────
_SAFE_PHOTO = re.compile(r'^[\w\-]{1,64}\.(jpg|jpeg|png|webp)$', re.IGNORECASE)
PHOTO_MAX_INPUT_BYTES = 20 * 1024 * 1024
PHOTO_MAX_DIMENSION = 1280
PHOTO_JPEG_QUALITY = 74

def compress_photo_bytes(raw):
    """Normalize an uploaded photo to an oriented, storage-efficient JPEG."""
    if not raw:
        raise ValueError('Image is empty')

    try:
        with Image.open(io.BytesIO(raw)) as source:
            source_format = (source.format or '').upper()
            source_orientation = source.getexif().get(274, 1)
            source.load()
            image = ImageOps.exif_transpose(source)

            # JPEG has no alpha channel. Composite transparent images over white.
            if image.mode in ('RGBA', 'LA') or 'transparency' in image.info:
                rgba = image.convert('RGBA')
                flattened = Image.new('RGB', rgba.size, 'white')
                flattened.paste(rgba, mask=rgba.getchannel('A'))
                image = flattened
            elif image.mode != 'RGB':
                image = image.convert('RGB')

            original_size = image.size
            image.thumbnail(
                (PHOTO_MAX_DIMENSION, PHOTO_MAX_DIMENSION),
                Image.Resampling.LANCZOS,
            )

            output = io.BytesIO()
            image.save(
                output,
                format='JPEG',
                quality=PHOTO_JPEG_QUALITY,
                optimize=True,
                progressive=True,
            )
            compressed = output.getvalue()

        # A browser-compressed JPEG may already be smaller than a second encode.
        # Keep it only when it is already compliant and needs no orientation fix.
        already_within_limit = max(original_size) <= PHOTO_MAX_DIMENSION
        if (source_format in ('JPEG', 'JPG') and source_orientation in (None, 1)
                and already_within_limit and len(raw) <= len(compressed)):
            return raw
        return compressed
    except (UnidentifiedImageError, OSError, ValueError) as exc:
        raise ValueError('Invalid or unsupported image file') from exc

def resolve_photos(d, username):
    """Replace photo_filename references with actual base64 img_data for PDF generation."""
    temp_dir = get_temp_photos_dir(username)
    for area in d.get('areas', []):
        for photo in area.get('photos', []):
            if not photo.get('img_data') and photo.get('photo_filename'):
                fname = photo['photo_filename']
                if _SAFE_PHOTO.match(fname):
                    fpath = os.path.join(temp_dir, fname)
                    if os.path.isfile(fpath):
                        try:
                            with open(fpath, 'rb') as f: raw = f.read()
                            ext = fname.rsplit('.', 1)[-1].lower()
                            mime = 'image/jpeg' if ext in ('jpg','jpeg') else f'image/{ext}'
                            photo['img_data'] = f'data:{mime};base64,{base64.b64encode(raw).decode()}'
                        except: pass
    return d

@app.route('/upload_temp_photo', methods=['POST'])
@login_required
def upload_temp_photo():
    username = session['username']
    temp_dir = get_temp_photos_dir(username)

    # ── Binary FormData upload (preferred path, ~33% smaller than base64 JSON) ──
    if 'photo' in request.files:
        try:
            f = request.files['photo']
            raw = f.read(PHOTO_MAX_INPUT_BYTES + 1)
            if len(raw) > PHOTO_MAX_INPUT_BYTES:
                return jsonify({'error': 'Image too large (max 20 MB)'}), 413
            compressed = compress_photo_bytes(raw)
            fname = f'{uuid.uuid4().hex}.jpg'
            fpath = os.path.join(temp_dir, fname)
            with open(fpath, 'wb') as fh: fh.write(compressed)
            return jsonify({
                'ok': True,
                'photo_filename': fname,
                'original_size': len(raw),
                'stored_size': len(compressed),
            })
        except ValueError as e:
            return jsonify({'error': str(e)}), 400
        except Exception as e:
            return jsonify({'error': str(e)}), 500

    # ── Legacy JSON base64 upload (fallback) ──
    data = request.json or {}
    img_data = data.get('img_data', '')
    if not img_data or ',' not in img_data:
        return jsonify({'error': 'No image data'}), 400
    try:
        header, b64 = img_data.split(',', 1)
        raw = base64.b64decode(b64)
        if len(raw) > PHOTO_MAX_INPUT_BYTES:
            return jsonify({'error': 'Image too large (max 20 MB)'}), 413
        compressed = compress_photo_bytes(raw)
        fname = f'{uuid.uuid4().hex}.jpg'
        fpath = os.path.join(temp_dir, fname)
        with open(fpath, 'wb') as fh: fh.write(compressed)
        return jsonify({
            'ok': True,
            'photo_filename': fname,
            'original_size': len(raw),
            'stored_size': len(compressed),
        })
    except ValueError as e:
        return jsonify({'error': str(e)}), 400
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/temp_photo/<filename>')
@login_required
def serve_temp_photo(filename):
    username = session['username']
    if not _SAFE_PHOTO.match(filename):
        return 'Invalid filename', 400
    fpath = os.path.join(get_temp_photos_dir(username), filename)
    if not os.path.isfile(fpath): return 'Not found', 404
    ext = filename.rsplit('.', 1)[-1].lower()
    mime = 'image/jpeg' if ext in ('jpg','jpeg') else f'image/{ext}'
    return send_file(fpath, mimetype=mime)

@app.route('/generate', methods=['POST'])
@login_required
def generate():
    username = session['username']
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict):
        return jsonify({'error': 'Invalid report data'}), 400

    try:
        d = resolve_photos(payload, username)
        cfg = load_config()
        date_str = _safe_report_filename_part(d.get('date'), 'Report').replace(' ', '_')
        day_str = _safe_report_filename_part(d.get('day_no'), 'Unnumbered')
        fname = f"Daily Report - PT GPA - KN - {date_str} (Day {day_str}).pdf"
        buf = generate_pdf(d, None, cfg)
    except Exception as e:
        app.logger.exception('PDF generation failed for user %s', username)
        return jsonify({'error': f'PDF generation failed: {e}'}), 500

    pdf_bytes = buf.getvalue()
    archive_failed = False
    try:
        archive_generated_report(username, fname, pdf_bytes, {
            'filename':     fname,
            'date':         d.get('date',''),
            'day_no':       d.get('day_no',''),
            'project_no':   d.get('project_no',''),
            'generated_at': datetime.now().strftime('%Y-%m-%d %H:%M'),
        })
    except Exception:
        # A temporary My Reports problem must not block a valid PDF download.
        archive_failed = True
        app.logger.exception('Could not archive generated PDF for user %s', username)

    buf.seek(0)
    response = send_file(
        buf,
        as_attachment=True,
        download_name=fname,
        mimetype='application/pdf',
    )
    response.headers['X-Report-Archive-Status'] = 'failed' if archive_failed else 'saved'
    return response

@app.route('/save_draft', methods=['POST'])
@login_required
def save_draft():
    username = session['username']
    data = request.json
    with open(get_draft_file(username),'w') as f:
        json.dump(data, f)
    save_draft_snapshot(username, data)
    log_activity(username, 'draft_saved', f"day={data.get('day_no','')} date={data.get('date','')}")
    ts = datetime.now().strftime('%H:%M')
    return jsonify({'ok':True, 'saved_at': ts})

@app.route('/preview', methods=['POST'])
@login_required
def preview():
    try:
        d = resolve_photos(request.json, session['username'])
        cfg = load_config()
        buf = generate_pdf(d, None, cfg)
        return send_file(buf, mimetype='application/pdf')
    except Exception as e:
        return jsonify({'error': f'PDF generation failed: {e}'}), 500

@app.route('/load_draft')
@login_required
def load_draft_route():
    df = get_draft_file(session['username'])
    if os.path.exists(df):
        with open(df) as f: return jsonify(json.load(f))
    return jsonify({})

@app.route('/get_config')
@login_required
def get_config():
    return jsonify(load_config())

@app.route('/save_config', methods=['POST'])
@login_required
def save_config_route():
    cfg = load_config()
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({'error':'Invalid settings data.'}), 400
    if 'projects' in data:
        if not session.get('is_admin', False):
            return jsonify({'error':'Only an admin can change Master Projects.'}), 403
        try:
            cfg['projects'] = normalize_projects(data['projects'], strict=True)
        except ValueError as exc:
            return jsonify({'error':str(exc)}), 400
    for k in ['company_name','customer','project_no','project_title','location','equipment',
              'prepared_by','checked_by','approved_by','areas','hours_options',
              'show_logo_gpa','show_logo_kn',
              'logo_gpa_w','logo_gpa_h','logo_gpa_y_off',
              'logo_kn_w','logo_kn_h','logo_kn_y_off',
              'manpower_db','ai_api_key',
              'site_name','pm_name','pm_title','letter_kn_attn','project_start_date']:
        if k in data: cfg[k] = data[k]
    if 'theme' in data:
        cfg['theme'].update(data['theme'])
    save_config(cfg)
    return jsonify({'ok':True})

@app.route('/upload_logo', methods=['POST'])
@login_required
def upload_logo():
    which = request.args.get('which','gpa')
    if 'file' not in request.files:
        return jsonify({'error':'no file'}), 400
    f = request.files['file']
    ext = os.path.splitext(f.filename)[1].lower() or '.png'
    fname = f'logo_{which}{ext}'
    save_path = os.path.join(LOGOS_DIR, fname)
    f.save(save_path)
    cfg = load_config()
    cfg[f'logo_{which}'] = save_path
    save_config(cfg)
    return jsonify({'ok':True, 'path':save_path})

@app.route('/remove_logo', methods=['POST'])
@login_required
def remove_logo():
    which = request.args.get('which','gpa')
    cfg = load_config()
    old = cfg.get(f'logo_{which}','')
    if old and os.path.isfile(old):
        try: os.remove(old)
        except: pass
    cfg[f'logo_{which}'] = ''
    save_config(cfg)
    return jsonify({'ok':True})

@app.route('/logo_status')
@login_required
def logo_status():
    cfg = load_config()
    resolved_gpa = resolve_logo_path(cfg.get('logo_gpa', ''), 'gpa')
    resolved_kn = resolve_logo_path(cfg.get('logo_kn', ''), 'kn')
    return jsonify({
        'gpa': bool(resolved_gpa),
        'kn': bool(resolved_kn),
        'gpa_source': 'uploaded' if _is_drawable_logo(cfg.get('logo_gpa', '')) else 'bundled',
        'kn_source': 'uploaded' if _is_drawable_logo(cfg.get('logo_kn', '')) else 'bundled',
    })



# ═══════════════════════════════════════════════════════════════════════════════
# Update 0.5 additions
# ═══════════════════════════════════════════════════════════════════════════════

ROMAN_MONTH = {1:'I',2:'II',3:'III',4:'IV',5:'V',6:'VI',
               7:'VII',8:'VIII',9:'IX',10:'X',11:'XI',12:'XII'}

# ── Letter storage helpers ────────────────────────────────────────────────────
def get_letters_dir(username):
    d = os.path.join(get_user_dir(username), 'letters')
    os.makedirs(d, exist_ok=True)
    return d

def get_letters_index(username):
    idx_file = os.path.join(get_letters_dir(username), 'index.json')
    if os.path.exists(idx_file):
        try:
            with open(idx_file) as f: return json.load(f)
        except: pass
    return []

def append_letter_index(username, entry):
    letters = get_letters_index(username)
    letters.insert(0, entry)
    idx_file = os.path.join(get_letters_dir(username), 'index.json')
    with open(idx_file,'w') as f: json.dump(letters, f, indent=2)

def next_letter_seq():
    cfg = load_config()
    seq = int(cfg.get('letter_seq_no', 1))
    cfg['letter_seq_no'] = seq + 1
    save_config(cfg)
    return seq

# ── Activity log ──────────────────────────────────────────────────────────────
def log_activity(username, action, details=''):
    log_file = ACTIVITY_LOG_FILE
    entry = {
        'ts':      datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        'user':    username,
        'action':  action,
        'details': details,
    }
    log = []
    if os.path.exists(log_file):
        try:
            with open(log_file) as f: log = json.load(f)
        except: pass
    log.insert(0, entry)
    log = log[:1000]
    with open(log_file,'w') as f: json.dump(log, f, indent=2)

# ── Versioned draft snapshot ──────────────────────────────────────────────────
_SNAPSHOT_FILENAME = re.compile(r'^\d{8}_\d{6}\.json$')

def save_draft_snapshot(username, data):
    snap_dir = os.path.join(get_user_dir(username), 'drafts')
    os.makedirs(snap_dir, exist_ok=True)
    ts = datetime.now().strftime('%Y%m%d_%H%M%S')
    snap_file = os.path.join(snap_dir, f'{ts}.json')
    with open(snap_file,'w') as f: json.dump(data, f)
    # Keep only last 20 snapshots
    snaps = sorted(
        name for name in os.listdir(snap_dir)
        if _SNAPSHOT_FILENAME.fullmatch(name)
    )
    for old in snaps[:-20]:
        try: os.remove(os.path.join(snap_dir, old))
        except: pass

def get_draft_snapshots(username):
    snap_dir = os.path.join(get_user_dir(username), 'drafts')
    if not os.path.isdir(snap_dir):
        return []
    snaps = sorted(os.listdir(snap_dir), reverse=True)
    result = []
    for s in snaps:
        if _SNAPSHOT_FILENAME.fullmatch(s):
            ts_raw = s.replace('.json','')
            try:
                ts_fmt = datetime.strptime(ts_raw, '%Y%m%d_%H%M%S').strftime('%d %b %Y %H:%M:%S')
            except:
                ts_fmt = ts_raw
            sz = round(os.path.getsize(os.path.join(snap_dir, s)) / 1024, 1)
            result.append({'filename': s, 'ts': ts_fmt, 'size_kb': sz})
    return result

# ── python-docx letter generation ─────────────────────────────────────────────
try:
    from docx import Document as DocxDocument
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_OK = True
except ImportError:
    DOCX_OK = False

def _letter_header(doc, d, cfg, seq_str, subject):
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    dt_str = d.get('date', datetime.now().strftime('%d %B %Y'))
    site   = cfg.get('site_name', 'Mangkajang')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(f' {site}, {dt_str}')
    run.font.size = Pt(11)
    doc.add_paragraph()
    for label, val in [('No', f': {seq_str}'), ('Perihal', f': {subject}')]:
        p = doc.add_paragraph()
        run = p.add_run(f'{label}\t{val}')
        run.font.size = Pt(11)
    doc.add_paragraph()
    p = doc.add_paragraph(); p.add_run('Yth.').font.size = Pt(11)
    recipient = d.get('recipient', 'PT Kertas Nusantara')
    p = doc.add_paragraph()
    run = p.add_run(recipient)
    run.font.size = Pt(11); run.bold = True
    if d.get('recipient_type', 'kn') == 'kn':
        p = doc.add_paragraph()
        p.add_run('Mangkajang, Tanjung Redeb, Berau Kalimantan Timur').font.size = Pt(11)
        attn = d.get('attn', cfg.get('letter_kn_attn', ''))
        p = doc.add_paragraph()
        p.add_run(f'Up: {attn}').font.size = Pt(11)
    else:
        p = doc.add_paragraph(); p.add_run(d.get('recipient_pos','')).font.size = Pt(11)
        p = doc.add_paragraph(); p.add_run('Di Tempat.').font.size = Pt(11)
    doc.add_paragraph()
    p = doc.add_paragraph(); p.add_run('Dengan Hormat,').font.size = Pt(11)
    doc.add_paragraph()

def _letter_closing(doc, cfg):
    p = doc.add_paragraph()
    p.add_run('Demikian yang Kami Sampaikan. Atas Perhatian dan kerjasamanya, Kami Ucapkan Terimakasih.').font.size = Pt(11)
    doc.add_paragraph()

def _letter_signoff_kn(doc, cfg):
    pm_name  = cfg.get('pm_name',  'Asril Naimy')
    pm_title = cfg.get('pm_title', 'Project Manager')
    table = doc.add_table(rows=4, cols=3)
    table.style = 'Table Grid'
    for i, h in enumerate(['PT. Garuda Prima Aksara','PT. Kertas Nusantara','PT. Kertas Nusantara']):
        c = table.cell(0, i); c.text = h
        c.paragraphs[0].runs[0].bold = True
        c.paragraphs[0].runs[0].font.size = Pt(10)
    table.cell(1,0).text = pm_name
    for c in [table.cell(2,i) for i in range(3)]:
        c.text = ''
    for i, pos in enumerate([pm_title,'Safety PT. Kertas Nusantara','Security PT. Kertas Nusantara']):
        c = table.cell(3,i); c.text = pos
        c.paragraphs[0].runs[0].font.size = Pt(10)

def _letter_signoff_internal(doc, cfg):
    pm_name  = cfg.get('pm_name',  'Asril Naimy')
    pm_title = cfg.get('pm_title', 'Project Manager')
    p = doc.add_paragraph(); p.add_run('Hormat kami,').font.size = Pt(11)
    for _ in range(3): doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run(f'    {pm_name}')
    run.font.size = Pt(11); run.underline = True
    p = doc.add_paragraph(); p.add_run(pm_title).font.size = Pt(11)

def _gen_letter_pekerja(doc, d, cfg, seq_str, subject):
    _letter_header(doc, d, cfg, seq_str, subject)
    body = d.get('body','Sehubungan adanya pengerjaan Electrical di PT. Garuda Prima Aksara Site PT. Kertas Nusantara. Oleh karena itu, kami meminta izin untuk masuk pekerja dan Induction sebagai berikut.')
    p = doc.add_paragraph(); p.add_run(body).font.size = Pt(11)
    doc.add_paragraph()
    workers = d.get('workers', [])
    table = doc.add_table(rows=1+len(workers), cols=2)
    table.style = 'Table Grid'
    for i, h in enumerate(['Nama','Posisi']):
        c = table.rows[0].cells[i]; c.text = h
        c.paragraphs[0].runs[0].bold = True; c.paragraphs[0].runs[0].font.size = Pt(11)
    for i, w in enumerate(workers):
        row = table.rows[i+1].cells
        row[0].text = w.get('name',''); row[1].text = w.get('position','')
    doc.add_paragraph()
    _letter_closing(doc, cfg)
    _letter_signoff_kn(doc, cfg)

def _gen_letter_alat_berat(doc, d, cfg, seq_str, subject):
    _letter_header(doc, d, cfg, seq_str, subject)
    body = d.get('body','Sehubungan akan dilakukannya pengerjaan Electrical, dengan ini kami mengajukan permohonan untuk meminjam alat berat kepada pihak PT. Kertas Nusantara')
    p = doc.add_paragraph(); p.add_run(body).font.size = Pt(11)
    doc.add_paragraph()
    for label, val in [('Hari/Tanggal', d.get('date_use','')),
                       ('Lokasi',        d.get('location_use','')),
                       ('Jenis Alat Berat', d.get('equipment_use',''))]:
        p = doc.add_paragraph()
        p.add_run(f'{label}\t\t: {val}').font.size = Pt(11)
    doc.add_paragraph()
    _letter_closing(doc, cfg)
    pm_name = cfg.get('pm_name','Asril Naimy')
    table = doc.add_table(rows=2, cols=2)
    table.style = 'Table Grid'
    table.cell(0,0).text = 'PT. Garuda Prima Aksara'
    table.cell(0,1).text = 'PT. Kertas Nusantara'
    table.cell(1,0).text = pm_name

def _gen_letter_sticker(doc, d, cfg, seq_str, subject):
    _letter_header(doc, d, cfg, seq_str, subject)
    body = d.get('body','Sehubung dengan adanya penambahan kendaraan dari PT. Garuda Prima Aksara Site PT. Kertas Nusantara. Oleh Karena itu, kami meminta identitas kendaraan (Sticker) sebagai persyarataan izin masuk, adapun jenis Kendaraan yang digunakan')
    p = doc.add_paragraph(); p.add_run(body).font.size = Pt(11)
    doc.add_paragraph()
    vehicles = d.get('vehicles', [])
    table = doc.add_table(rows=1+len(vehicles), cols=3)
    table.style = 'Table Grid'
    for i, h in enumerate(['NO','JENIS KENDARAAN','NOMOR POLISI']):
        c = table.rows[0].cells[i]; c.text = h
        c.paragraphs[0].runs[0].bold = True
    for i, v in enumerate(vehicles):
        row = table.rows[i+1].cells
        row[0].text = str(i+1); row[1].text = v.get('type',''); row[2].text = v.get('plate','')
    doc.add_paragraph()
    _letter_closing(doc, cfg)
    _letter_signoff_kn(doc, cfg)

def _gen_letter_umum(doc, d, cfg, seq_str, subject):
    _letter_header(doc, d, cfg, seq_str, subject)
    p = doc.add_paragraph(); p.add_run(d.get('body','')).font.size = Pt(11)
    doc.add_paragraph()
    _letter_closing(doc, cfg)
    if d.get('recipient_type','kn') == 'kn':
        _letter_signoff_kn(doc, cfg)
    else:
        _letter_signoff_internal(doc, cfg)

def _gen_letter_sp(doc, d, cfg, seq_str):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    sp_level = str(d.get('sp_level','3'))
    sp_label = {'1':'Pertama','2':'Kedua','3':'Ketiga'}.get(sp_level,'Ketiga')
    dt_str   = d.get('date', datetime.now().strftime('%d %B %Y'))
    site     = cfg.get('site_name','Mangkajang')
    emp_name = d.get('employee_name','')
    emp_pos  = d.get('employee_position','')
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run(f' {site}, {dt_str}').font.size = Pt(11)
    doc.add_paragraph()
    for label, val in [('No', f': {seq_str}'), ('Perihal', f': Peringatan ke-{sp_level}')]:
        p = doc.add_paragraph(); p.add_run(f'{label}\t{val}').font.size = Pt(11)
    doc.add_paragraph()
    p = doc.add_paragraph(); p.add_run('Yth.').font.size = Pt(11)
    p = doc.add_paragraph(); p.add_run(f'Sdr. {emp_name}').font.size = Pt(11)
    p = doc.add_paragraph(); p.add_run(emp_pos).font.size = Pt(11)
    p = doc.add_paragraph(); p.add_run('Di Tempat.').font.size = Pt(11)
    doc.add_paragraph()
    p = doc.add_paragraph(); p.add_run('Dengan hormat,').font.size = Pt(11)
    doc.add_paragraph()
    body = d.get('body','Sehubungan dengan evaluasi kinerja dan kedisiplinan kerja, kami mencatat bahwa Saudara telah beberapa kali tidak hadir tanpa keterangan yang jelas dan dapat dipertanggungjawabkan.')
    p = doc.add_paragraph(); p.add_run(body).font.size = Pt(11)
    doc.add_paragraph()
    if sp_level == '3':
        p = doc.add_paragraph(); p.add_run('Sebelumnya, perusahaan telah memberikan:').font.size = Pt(11)
        p = doc.add_paragraph(); p.add_run('Peringatan Pertama (Lisan)').font.size = Pt(11)
        p = doc.add_paragraph(); p.add_run('Peringatan Kedua (Lisan)').font.size = Pt(11)
        doc.add_paragraph()
    consequence = d.get('consequence','Apabila Saudara kembali tidak masuk kerja tanpa keterangan, perusahaan akan melakukan pemutusan hubungan kerja (PHK).')
    p = doc.add_paragraph()
    p.add_run(f'Oleh karena itu, perusahaan memberikan Surat Peringatan {sp_label} (SP{sp_level}) sebagai peringatan. {consequence}').font.size = Pt(11)
    doc.add_paragraph()
    p = doc.add_paragraph(); p.add_run('Demikian surat peringatan ini disampaikan untuk menjadi perhatian dan dilaksanakan sebagaimana mestinya.').font.size = Pt(11)
    doc.add_paragraph()
    _letter_signoff_internal(doc, cfg)

def _gen_letter_phk(doc, d, cfg, seq_str):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    dt_str   = d.get('date', datetime.now().strftime('%d %B %Y'))
    site     = cfg.get('site_name','Mangkajang')
    emp_name = d.get('employee_name','')
    emp_pos  = d.get('employee_position','')
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run(f' {site}, {dt_str}').font.size = Pt(11)
    doc.add_paragraph()
    for label, val in [('No', f': {seq_str}'), ('Perihal', ': Surat Pemutusan Hubungan Kerja (PHK)')]:
        p = doc.add_paragraph(); p.add_run(f'{label}\t{val}').font.size = Pt(11)
    doc.add_paragraph()
    p = doc.add_paragraph(); p.add_run('Yth.').font.size = Pt(11)
    p = doc.add_paragraph(); p.add_run(f'Sdr. {emp_name}').font.size = Pt(11)
    p = doc.add_paragraph(); p.add_run(emp_pos).font.size = Pt(11)
    p = doc.add_paragraph(); p.add_run('Di Tempat.').font.size = Pt(11)
    doc.add_paragraph()
    p = doc.add_paragraph(); p.add_run('Dengan hormat,').font.size = Pt(11)
    doc.add_paragraph()
    sp_ref = d.get('sp_reference','Surat Peringatan Ketiga (SP 3)')
    body   = d.get('body',f'Berdasarkan {sp_ref} yang telah diberikan kepada Saudara, dimana telah disepakati bahwa apabila Saudara kembali tidak masuk kerja tanpa keterangan, maka akan dilakukan Pemutusan Hubungan Kerja (PHK).')
    p = doc.add_paragraph(); p.add_run(body).font.size = Pt(11)
    doc.add_paragraph()
    eff_date = d.get('effective_date', dt_str)
    p = doc.add_paragraph()
    p.add_run(f'Sehubungan dengan hal tersebut, perusahaan dengan ini memutuskan untuk melakukan Pemutusan Hubungan Kerja (PHK) terhadap Saudara {emp_name}, terhitung sejak tanggal {eff_date}.').font.size = Pt(11)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run('Kami mengucapkan terima kasih atas kontribusi yang telah Saudara berikan selama bekerja di perusahaan. Segala hak dan kewajiban yang timbul akibat keputusan ini akan diselesaikan sesuai dengan ketentuan yang berlaku di perusahaan.').font.size = Pt(11)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run('Demikian surat ini dibuat untuk dapat dipahami dan dilaksanakan sebagaimana mestinya.').font.size = Pt(11)
    doc.add_paragraph()
    _letter_signoff_internal(doc, cfg)

def generate_letter_docx(d, cfg):
    from io import BytesIO
    letter_type = d.get('letter_type','umum')
    seq  = next_letter_seq()
    dt   = datetime.now()
    seq_str = f"{seq:04d}/GPA-KN/{ROMAN_MONTH[dt.month]}/{dt.year}"
    doc = DocxDocument()
    for section in doc.sections:
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(3)
        section.right_margin  = Cm(2.5)
    doc.styles['Normal'].font.name = 'Times New Roman'
    doc.styles['Normal'].font.size = Pt(11)
    subject = d.get('subject','Permohonan')
    dispatch = {
        'pekerja':   _gen_letter_pekerja,
        'alat_berat': _gen_letter_alat_berat,
        'sticker':   _gen_letter_sticker,
        'sp':        lambda dc, dd, cfg2, ss, sub: _gen_letter_sp(dc, dd, cfg2, ss),
        'phk':       lambda dc, dd, cfg2, ss, sub: _gen_letter_phk(dc, dd, cfg2, ss),
    }
    fn = dispatch.get(letter_type, _gen_letter_umum)
    fn(doc, d, cfg, seq_str, subject)
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf, seq_str

# ── Letter routes ─────────────────────────────────────────────────────────────
@app.route('/letters')
@login_required
def letters_page():
    username = session['username']
    cfg      = load_config()
    letters  = get_letters_index(username)
    return render_template('letters.html',
        username=username,
        is_admin=session.get('is_admin',False),
        letters=letters,
        cfg=cfg,
    )

@app.route('/letters/generate', methods=['POST'])
@login_required
def generate_letter():
    if not DOCX_OK:
        return jsonify({'error':'python-docx not installed. Run: pip install python-docx'}), 500
    username = session['username']
    d        = request.json
    cfg      = load_config()
    d.setdefault('date', datetime.now().strftime('%d %B %Y'))
    try:
        buf, seq_str = generate_letter_docx(d, cfg)
    except Exception as e:
        return jsonify({'error': str(e)}), 500
    letter_type = d.get('letter_type','umum')
    type_labels = {
        'pekerja':   'Permohonan Masuk Pekerja',
        'alat_berat':'Permohonan Alat Berat',
        'sticker':   'Permohonan Sticker',
        'sp':        f"SP{d.get('sp_level','3')}",
        'phk':       'PHK',
        'umum':      'Surat Umum',
    }
    type_label = type_labels.get(letter_type, letter_type)
    subject    = d.get('subject', type_label)
    safe_seq   = seq_str.replace('/','-')
    fname      = f"{safe_seq}_{type_label}.docx"
    ldir  = get_letters_dir(username)
    fpath = os.path.join(ldir, fname)
    with open(fpath,'wb') as f: f.write(buf.getvalue())
    append_letter_index(username, {
        'filename':     fname,
        'type':         letter_type,
        'type_label':   type_label,
        'subject':      subject,
        'seq_str':      seq_str,
        'date':         d.get('date',''),
        'generated_at': datetime.now().strftime('%Y-%m-%d %H:%M'),
        'size_kb':      round(os.path.getsize(fpath)/1024, 1),
    })
    log_activity(username, 'letter_generated', f'{type_label} - {seq_str}')
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=fname,
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

@app.route('/letters/download/<path:filename>')
@login_required
def download_letter(filename):
    username = session['username']
    fpath = os.path.join(get_letters_dir(username), filename)
    if not os.path.isfile(fpath): return 'Not found', 404
    return send_file(fpath, as_attachment=True, download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

@app.route('/letters/delete', methods=['POST'])
@login_required
def delete_letter():
    username = session['username']
    fname = request.json.get('filename','')
    fpath = os.path.join(get_letters_dir(username), fname)
    if os.path.isfile(fpath): os.remove(fpath)
    letters = [l for l in get_letters_index(username) if l.get('filename') != fname]
    with open(os.path.join(get_letters_dir(username),'index.json'),'w') as f:
        json.dump(letters, f, indent=2)
    return jsonify({'ok':True})

# ── Yesterday crew ────────────────────────────────────────────────────────────
@app.route('/load_yesterday')
@login_required
def load_yesterday():
    username  = session['username']
    draft_file = get_draft_file(username)
    if os.path.exists(draft_file):
        try:
            with open(draft_file) as f: draft = json.load(f)
            return jsonify({
                'ok': True,
                'indirect_manpower': draft.get('indirect_manpower',[]),
                'areas': [{'id':a.get('id'),
                           'manpower':a.get('manpower',[]),
                           'indirect_manpower':a.get('indirect_manpower',[])}
                          for a in draft.get('areas',[])],
            })
        except: pass
    return jsonify({'ok':False,'reason':'No draft found'})

# ── Draft snapshots ───────────────────────────────────────────────────────────
@app.route('/draft_snapshots')
@login_required
def draft_snapshots():
    return jsonify(get_draft_snapshots(session['username']))

@app.route('/draft_snapshots/load/<filename>')
@login_required
def load_draft_snapshot(filename):
    if not _SNAPSHOT_FILENAME.fullmatch(filename):
        return jsonify({'error':'Not found'}), 404
    username = session['username']
    snap_dir = os.path.join(get_user_dir(username),'drafts')
    fpath    = os.path.join(snap_dir, filename)
    if not os.path.isfile(fpath): return jsonify({'error':'Not found'}), 404
    with open(fpath) as f: data = json.load(f)
    return jsonify(data)

# ── Activity log (admin) ──────────────────────────────────────────────────────
@app.route('/admin/activity_log')
@admin_required
def activity_log_page():
    log_file = ACTIVITY_LOG_FILE
    log = []
    if os.path.exists(log_file):
        try:
            with open(log_file) as f: log = json.load(f)
        except: pass
    return render_template('activity_log.html',
        username=session['username'],
        is_admin=True,
        log=log,
    )

# ── AI Chat Assistant ────────────────────────────────────────────────────────
import anthropic as _anthropic_mod

_AI_SYSTEM_PROMPT_TEMPLATE = string.Template("""You are a daily report assistant for PT. Garuda Prima Aksara (electrical construction project).
The user will describe their day's work in natural language (often in Indonesian/Bahasa).
Your job is to extract structured data and fill the daily report form.

The report has these sections:
1. **Report Info**: date, day_no (auto-calc from project_start_date), project_no, location, customer, equipment, project_title, photo_documentation_title, prepared_by, checked_by, approved_by
2. **Weather**: Morning, Afternoon, Evening (Cerah/Berawan/Hujan/Gerimis), Wind (Calm/Light/Moderate/Strong), Temperature (free text e.g. "30°C"), Impact (None/Minor/Stopped Work)
3. **Global Indirect Manpower**: list of {name, role, hours}
4. **Overall Progress**: list of {description, duration, weight_factor, start, finish, cumulative_previous_plan, cumulative_previous_actual, this_period_plan, this_period_actual, cumulative_to_date_plan, cumulative_to_date_actual}
5. **Areas** (can be multiple): each has:
   - id (area name like MA-14, MA-23, etc.)
   - activities_today: list of strings
   - activities_tomorrow: list of strings
   - manpower: list of {name, role, task, hours}
   - indirect_manpower: list of {name, role, hours}
   - constraints: string
   - remarks: string
6. **Global Remarks**: free text
7. **Sign-offs**: auto-filled from config

Available areas: $areas
Available manpower (name to role): $manpower
Config defaults: project_no=$project_no, location=$location, customer=$customer, project_title=$project_title, prepared_by=$prepared_by, checked_by=$checked_by, approved_by=$approved_by
Project start date: $project_start_date

RULES:
- Respond in the same language the user uses (Indonesian or English).
- When the user describes work, extract as much data as possible into the form fields.
- For manpower names, fuzzy-match against the known list (e.g. "Budi" -> closest match in list).
- Always return a JSON block with your form updates AND a text reply.
- If info is missing, ask about it naturally - don't interrogate field by field.
- You can suggest using yesterday's crew if the user hasn't specified manpower.
- Default hours to "07:00 - 17:00" if not specified.
- When you have enough data for a complete report, tell the user they can review the form and generate.

Your response MUST be valid JSON with exactly this structure:
{
  "reply": "Your conversational reply to the user (markdown OK)",
  "updates": {
    "date": "...",
    "photo_documentation_title": "...",
    "weather": {"Morning": "...", ...},
    "overall_progress": [{"description": "...", "weight_factor": "...", "cumulative_to_date_plan": "...", "cumulative_to_date_actual": "..."}],
    "areas": [{"id": "MA-14", "activities_today": [...], "activities_tomorrow": [...], "manpower": [...], "constraints": "...", "remarks": ""}]
  },
  "missing": ["list of important fields still missing"],
  "ready": false
}

Only include fields in "updates" that you are updating. Omit fields you don't have data for.
Set "ready": true when you believe the report has enough data to generate (at minimum: date, weather, at least 1 area with activities).
""")

@app.route('/ai/chat', methods=['POST'])
@login_required
def ai_chat():
    cfg = load_config()
    api_key = os.environ.get('ANTHROPIC_API_KEY', cfg.get('ai_api_key', '')).strip()
    if not api_key:
        return jsonify({'error': 'No AI API key configured. Go to Settings → AI Assistant to add your Anthropic API key.'}), 400

    data = request.json or {}
    user_msg = (data.get('message') or '').strip()
    history = data.get('history') or []
    current_form = data.get('current_form') or {}

    if not user_msg:
        return jsonify({'error': 'Empty message'}), 400

    mp_db = cfg.get('manpower_db', MANPOWER_DB)
    mp_summary = ', '.join(f"{m['name']} ({m['role']})" for m in mp_db[:40])
    if len(mp_db) > 40:
        mp_summary += f' ... and {len(mp_db)-40} more'

    system = _AI_SYSTEM_PROMPT_TEMPLATE.safe_substitute(
        areas=', '.join(cfg.get('areas', AREA_LIST)),
        manpower=mp_summary,
        project_no=cfg.get('project_no', ''),
        location=cfg.get('location', ''),
        customer=cfg.get('customer', ''),
        project_title=cfg.get('project_title', ''),
        prepared_by=cfg.get('prepared_by', ''),
        checked_by=cfg.get('checked_by', ''),
        approved_by=cfg.get('approved_by', ''),
        project_start_date=cfg.get('project_start_date', ''),
    )

    if current_form:
        system += f"\n\nCurrent form state (what's already filled):\n{json.dumps(current_form, ensure_ascii=False, default=str)[:3000]}"

    messages = []
    for h in history[-20:]:
        messages.append({"role": h.get("role", "user"), "content": h.get("content", "")})
    messages.append({"role": "user", "content": user_msg})

    try:
        client = _anthropic_mod.Anthropic(api_key=api_key)
        resp = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=2048,
            system=system,
            messages=messages,
        )
        raw = resp.content[0].text.strip()

        # Extract JSON from response (handle markdown code blocks)
        if '```json' in raw:
            raw = raw.split('```json', 1)[1].split('```', 1)[0].strip()
        elif '```' in raw:
            raw = raw.split('```', 1)[1].split('```', 1)[0].strip()

        result = json.loads(raw)
        return jsonify(result)
    except json.JSONDecodeError:
        return jsonify({
            'reply': raw if 'raw' in dir() else 'Sorry, I had trouble formatting my response.',
            'updates': {},
            'missing': [],
            'ready': False,
        })
    except _anthropic_mod.APIError as e:
        return jsonify({'error': f'AI API error: {str(e)}'}), 500
    except Exception as e:
        return jsonify({'error': f'AI error: {str(e)}'}), 500


# ── Field submissions (supervisor mobile app) ─────────────────────────────────

def get_field_submissions_dir(username):
    d = os.path.join(get_user_dir(username), 'field_submissions')
    os.makedirs(d, exist_ok=True)
    return d

def load_field_submissions(username, date_str):
    path = os.path.join(get_field_submissions_dir(username), f"{date_str}.json")
    if os.path.exists(path):
        try:
            with open(path) as f: return json.load(f)
        except: pass
    return []

def save_field_submissions(username, date_str, submissions):
    d = get_field_submissions_dir(username)
    os.makedirs(d, exist_ok=True)
    with open(os.path.join(d, f"{date_str}.json"), 'w') as f:
        json.dump(submissions, f, ensure_ascii=False, indent=2)

@app.route('/field')
@login_required
def field_page():
    cfg = load_config()
    return render_template('field.html', config=cfg)

@app.route('/field/load')
@login_required
def field_load():
    username = session['username']
    date_str = request.args.get('date', datetime.now().strftime('%Y-%m-%d'))
    return jsonify(load_field_submissions(username, date_str))

@app.route('/field/submit', methods=['POST'])
@login_required
def field_submit():
    username = session['username']
    data = request.json or {}
    date_str = data.get('date', datetime.now().strftime('%Y-%m-%d'))
    area_id  = (data.get('area_id') or '').strip()
    if not area_id:
        return jsonify({'error': 'area_id required'}), 400

    submissions = load_field_submissions(username, date_str)
    # Upsert: replace existing entry for same area
    submissions = [s for s in submissions if s.get('area_id') != area_id]
    submissions.append({
        'username':           username,
        'submitted_at':       datetime.now().strftime('%Y-%m-%d %H:%M'),
        'date':               date_str,
        'area_id':            area_id,
        'activities_today':   data.get('activities_today', []),
        'activities_tomorrow':data.get('activities_tomorrow', []),
        'constraints':        data.get('constraints', ''),
        'manpower':           data.get('manpower', []),
        'photos':             data.get('photos', []),
    })
    save_field_submissions(username, date_str, submissions)
    log_activity(username, 'field_submitted', f"area={area_id} date={date_str}")
    return jsonify({'ok': True})

@app.route('/admin/all_submissions')
@admin_required
def admin_all_submissions():
    date_str = request.args.get('date', datetime.now().strftime('%Y-%m-%d'))
    users = load_users()
    all_subs = []
    for uname in users:
        subs = load_field_submissions(uname, date_str)
        all_subs.extend(subs)
    # Sort by submitted_at
    all_subs.sort(key=lambda s: s.get('submitted_at', ''))
    return jsonify({'date': date_str, 'submissions': all_subs})

@app.route('/admin/merge_submissions', methods=['POST'])
@admin_required
def admin_merge_submissions():
    username = session['username']
    data = request.json or {}
    incoming = data.get('submissions', [])
    if not incoming:
        return jsonify({'error': 'No submissions provided'}), 400

    # Load existing draft (or empty)
    df = get_draft_file(username)
    draft = {}
    if os.path.exists(df):
        try:
            with open(df) as f: draft = json.load(f)
        except: pass

    existing_areas = {a.get('id'): a for a in draft.get('areas', [])}

    for sub in incoming:
        area_id = sub.get('area_id', '')
        area_block = {
            'id':                  area_id,
            'activities_today':    sub.get('activities_today', []),
            'activities_tomorrow': sub.get('activities_tomorrow', []),
            'constraints':         sub.get('constraints', ''),
            'manpower':            sub.get('manpower', []),
            'indirect_manpower':   [],
            'photos':              sub.get('photos', []),
            'remarks':             '',
        }
        existing_areas[area_id] = area_block

    draft['areas'] = list(existing_areas.values())
    if not draft.get('date') and incoming:
        draft['date'] = incoming[0].get('date', '')

    with open(df, 'w') as f:
        json.dump(draft, f, ensure_ascii=False, indent=2)
    save_draft_snapshot(username, draft)
    log_activity(username, 'submissions_merged', f"areas={[s.get('area_id') for s in incoming]}")
    return jsonify({'ok': True, 'merged': len(incoming)})


if __name__ == '__main__':
    import webbrowser, threading, time, socket

    try:
        local_ip = socket.gethostbyname(socket.gethostname())
    except Exception:
        local_ip = 'your-pc-ip'

    print()
    print("  ============================================")
    print("   Daily Report App v2  --  PT. GPA")
    print("  ============================================")
    print(f"  Local  : http://localhost:5050")
    print(f"  Network: http://{local_ip}:5050")
    print("  Browser opening automatically...")
    print("  To stop: close this window or press Ctrl+C")
    print()

    def _open_browser():
        time.sleep(1.5)
        webbrowser.open('http://localhost:5050')

    threading.Thread(target=_open_browser, daemon=True).start()
    app.run(host='0.0.0.0', port=5050, debug=False)
