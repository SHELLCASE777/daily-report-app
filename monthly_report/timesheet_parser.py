from __future__ import annotations

import io
from collections import defaultdict
from datetime import date, datetime
from typing import Any


def parse_timesheet(fileobj) -> dict[str, Any]:
    """
    Parse an Excel (.xlsx) timesheet and return aggregated manpower figures.

    Supports two layouts:
      A — Employee rows × date columns (columns headed 1–31 or date values).
          Each numeric/mark cell means the employee worked that day.
      B — Daily summary rows: (date | headcount | man-hours).

    Returns:
        {
          "total_manpower": int,       # peak daily headcount
          "total_man_hours": float,    # sum of all hours
          "daily_breakdown": [{"date": str, "headcount": int, "hours": float}, ...]
        }
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to parse Excel timesheets") from exc

    raw = fileobj.read() if hasattr(fileobj, "read") else fileobj
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    rows = [list(row) for row in ws.iter_rows(values_only=True)]
    wb.close()

    rows = [r for r in rows if any(c is not None for c in r)]
    if not rows:
        raise ValueError("Timesheet appears to be empty.")

    result = _try_layout_a(rows) or _try_layout_b(rows)
    if not result:
        raise ValueError(
            "Could not detect a recognised timesheet format. "
            "Expected either (A) an employee × date matrix with day-number column headers (1–31), "
            "or (B) daily summary rows: date | headcount | man-hours."
        )
    return result


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

_ATTENDANCE_MARKS = {"v", "✓", "√", "x", "y", "h", "hadir", "p", "1"}
_SKIP_LABELS = {"total", "jumlah", "amount", "sub total", "sub-total", "grand total"}


def _num(cell, default: float = 0.0) -> float:
    if isinstance(cell, bool):
        return default
    if isinstance(cell, (int, float)):
        return float(cell)
    try:
        return float(str(cell or "").replace(",", "").strip())
    except ValueError:
        return default


def _try_layout_a(rows: list[list]) -> dict[str, Any] | None:
    """Employee × date matrix — detect header row containing day numbers 1–31."""
    for header_idx, row in enumerate(rows[:20]):
        day_cols: dict[int, int] = {}  # column_index -> day_number (1-31)
        for j, cell in enumerate(row):
            if isinstance(cell, bool):
                continue
            if isinstance(cell, int) and 1 <= cell <= 31:
                day_cols[j] = cell
            elif isinstance(cell, (date, datetime)):
                day_cols[j] = cell.day if isinstance(cell, datetime) else cell.day
        if len(day_cols) < 10:
            continue

        # Found a header row with enough day columns
        daily: dict[int, dict[str, float]] = defaultdict(lambda: {"headcount": 0.0, "hours": 0.0})

        for data_row in rows[header_idx + 1 :]:
            if not any(data_row):
                continue
            first = str(data_row[0] or "").strip().lower()
            if first in _SKIP_LABELS:
                continue
            for col_idx, day_num in day_cols.items():
                if col_idx >= len(data_row):
                    continue
                cell = data_row[col_idx]
                mark = str(cell or "").strip().lower()
                if mark in _ATTENDANCE_MARKS:
                    daily[day_num]["headcount"] += 1
                    daily[day_num]["hours"] += 8.0
                else:
                    h = _num(cell)
                    if h > 0:
                        daily[day_num]["headcount"] += 1
                        daily[day_num]["hours"] += h if h > 1 else 8.0

        breakdown = [
            {"date": f"day-{d:02d}", "headcount": int(daily[d]["headcount"]), "hours": daily[d]["hours"]}
            for d in sorted(daily)
            if daily[d]["headcount"] > 0
        ]
        if not breakdown:
            return None

        return {
            "total_manpower": max(d["headcount"] for d in breakdown),
            "total_man_hours": round(sum(d["hours"] for d in breakdown), 2),
            "daily_breakdown": breakdown,
        }
    return None


def _try_layout_b(rows: list[list]) -> dict[str, Any] | None:
    """Daily summary rows: first cell is a date/day, subsequent cells are headcount and hours."""
    breakdown = []
    for row in rows:
        cells = [c for c in row if c is not None]
        if len(cells) < 2:
            continue

        date_cell = cells[0]
        if isinstance(date_cell, (date, datetime)):
            date_str = (
                date_cell.strftime("%Y-%m-%d")
                if isinstance(date_cell, datetime)
                else date_cell.strftime("%Y-%m-%d")
            )
            day_num = date_cell.day if isinstance(date_cell, datetime) else date_cell.day
        elif isinstance(date_cell, (int, float)) and 1 <= int(date_cell) <= 31:
            day_num = int(date_cell)
            date_str = f"day-{day_num:02d}"
        else:
            s = str(date_cell or "").strip()
            try:
                sep = "/" if "/" in s else "-"
                parts = s.split(sep)
                day_num = int(parts[0]) if len(parts) >= 2 else int(parts[0])
                if not 1 <= day_num <= 31:
                    continue
                date_str = f"day-{day_num:02d}"
            except (ValueError, IndexError):
                continue

        nums = [_num(c) for c in cells[1:] if _num(c, -1) >= 0]
        if not nums:
            continue
        headcount = int(round(nums[0]))
        hours = nums[1] if len(nums) > 1 else headcount * 8.0
        if headcount <= 0 and hours <= 0:
            continue
        breakdown.append({"date": date_str, "headcount": headcount, "hours": hours})

    if len(breakdown) < 3:
        return None

    return {
        "total_manpower": max(d["headcount"] for d in breakdown),
        "total_man_hours": round(sum(d["hours"] for d in breakdown), 2),
        "daily_breakdown": breakdown,
    }
